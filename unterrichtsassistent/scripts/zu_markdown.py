#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Wandelt Dateien aus eingang/ in Markdown um (KONZEPT.md Kap. 5.1, 6.2).

Was das Skript macht - immer gleich, deshalb Skript (Designprinzip 3):
  - Text, Ueberschriften, Tabellen, Aufzaehlungen aus Word, PDF,
    PowerPoint, Excel und Textdateien holen
  - Bilder herausloesen und an der richtigen Stelle verlinken
  - Schreiblinien als [Schreibzeilen: n] erkennen (Kap. 5.4)
  - einen Metadatenkopf anlegen, in dem die feststehenden Felder gefuellt
    sind (schema, original, status, grafiken, Datum)

Was es nicht macht - das entscheidet Claude beim Lesen:
  - Fach, Klasse, Thema, Typ, Beschreibung, Herkunft, Quelle
  - Bilder ansehen und beschreiben, Deko streichen
  - Scans und Fotos lesen

Jede Datei landet in einem eigenen Arbeitsbereich unter
_system/aufnahme/<name>/ mit inhalt.md, bilder/ und auftrag.json. Dort
bearbeitet Claude sie, danach legt ablegen.py sie in die Wissensbasis.

Die Layout-Treue ist ausdruecklich kein Ziel: Das Original bleibt in
originale/ erhalten. Die .md ist die inhaltliche Abschrift, aus der spaeter
neues Material entsteht.

Aufruf:
    python zu_markdown.py --ordner "<Arbeitsordner>" --anzahl 10
    python zu_markdown.py --datei blatt.docx --ziel ausgabe/
"""

import argparse
import json
import logging
import math
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import gemeinsam as g                      # noqa: E402
import bilder_extrahieren as bx            # noqa: E402

logging.getLogger("pypdf").setLevel(logging.ERROR)

STANDARD_PORTION = 10

BILDDATEIEN = (".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic")
UMWEG_WORD = (".doc", ".odt", ".rtf", ".pages")
UMWEG_PPTX = (".ppt", ".odp", ".key")
UMWEG_XLSX = (".xls", ".ods", ".numbers")
IGNORIEREN = ("desktop.ini", "thumbs.db", ".ds_store")

# Ueberschriften, hinter denen vermutlich ein Loesungsteil beginnt
LOESUNGSTEIL = re.compile(
    r"(teacher'?s key|answer key|answers|solutions?|lösung|loesung|"
    r"erwartungshorizont|musterlösung|lehrerexemplar)", re.IGNORECASE)

SCHREIBLINIE = re.compile(r"^[\s_\-–.]*_{8,}[\s_\-–.]*$")


# --------------------------------------------------------------------------
# Hilfsmittel
# --------------------------------------------------------------------------

def schreibzeilen_zusammenfassen(zeilen):
    """Folgen von Schreiblinien werden zu [Schreibzeilen: n]."""
    ergebnis, anzahl = [], 0
    for zeile in zeilen + [None]:
        if zeile is not None and SCHREIBLINIE.match(zeile):
            anzahl += 1
            continue
        if anzahl:
            ergebnis.append("[Schreibzeilen: %d]" % anzahl)
            anzahl = 0
        if zeile is not None:
            ergebnis.append(zeile)
    return ergebnis


def tabellenzelle(text):
    text = (text or "").strip().replace("|", "\\|")
    return re.sub(r"\s*\n\s*", "<br>", text)


def pipe_tabelle(zeilen):
    """Liste von Zeilen (Listen von Zellen) als Markdown-Tabelle."""
    zeilen = [z for z in zeilen if any((c or "").strip() for c in z)]
    if not zeilen:
        return ""
    breite = max(len(z) for z in zeilen)
    zeilen = [list(z) + [""] * (breite - len(z)) for z in zeilen]
    aus = ["| " + " | ".join(tabellenzelle(c) for c in zeilen[0]) + " |",
           "| " + " | ".join("---" for _ in range(breite)) + " |"]
    for zeile in zeilen[1:]:
        aus.append("| " + " | ".join(tabellenzelle(c) for c in zeile) + " |")
    return "\n".join(aus)


def aufraeumen(text):
    """Hoechstens eine Leerzeile am Stueck, sauberes Ende."""
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def bildzeile(name, hinweis=""):
    zusatz = (" <!-- %s -->" % hinweis) if hinweis else ""
    return "![Bild: beschreiben oder streichen](bilder/%s)%s" % (name, zusatz)


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

def _stil_ebene(stil):
    """Ueberschriftenebene einer Formatvorlage: 1 = #, 2 = ## ..., 0 = keine,
    -1 = Untertitel. Sucht erst nach Word-Standardnamen (auch in der
    Vererbungskette), dann nach der Gliederungsebene, zuletzt nach dem Namen -
    Schulvorlagen heissen oft 'VTitle' oder 'AB-Ueberschrift'."""
    geprueft = 0
    aktuell = stil
    while aktuell is not None and geprueft < 6:
        name = (aktuell.name or "").strip()
        klein = name.lower()
        if klein in ("title", "titel"):
            return 1
        if klein in ("subtitle", "untertitel"):
            return -1
        treffer = re.match(r"(heading|überschrift|ueberschrift)\s*(\d)", klein)
        if treffer:
            return min(int(treffer.group(2)) + 1, 6)
        try:
            gliederung = aktuell.element.pPr.find(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl")
            if gliederung is not None:
                wert = int(gliederung.get(
                    "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"))
                if wert < 9:
                    return min(wert + 2, 6)
        except Exception:
            pass
        aktuell = getattr(aktuell, "base_style", None)
        geprueft += 1
    klein = (stil.name or "").lower()
    if re.search(r"sub|unter", klein) and re.search(r"title|titel|head|überschrift", klein):
        return -1
    if re.search(r"title|titel", klein):
        return 1
    if re.search(r"head|überschrift|ueberschrift", klein):
        return 2
    return 0


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _wert(element):
    return element.get(W + "val") if element is not None else None


def _roemisch(zahl):
    ergebnis = ""
    for wert, zeichen in ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
                          (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
                          (5, "v"), (4, "iv"), (1, "i")):
        while zahl >= wert:
            ergebnis += zeichen
            zahl -= wert
    return ergebnis


def _zahl_formatieren(zahl, format_):
    if format_ == "lowerLetter":
        return chr(96 + ((zahl - 1) % 26) + 1)
    if format_ == "upperLetter":
        return chr(64 + ((zahl - 1) % 26) + 1)
    if format_ == "lowerRoman":
        return _roemisch(zahl)
    if format_ == "upperRoman":
        return _roemisch(zahl).upper()
    if format_ == "decimalZero":
        return "%02d" % zahl
    return str(zahl)


class Nummerierung:
    """Rechnet Words automatische Nummerierung nach.

    Word speichert bei nummerierten Absaetzen nicht "1.", sondern nur einen
    Verweis auf eine Nummerierungsdefinition. Die Nummer selbst entsteht erst
    beim Anzeigen. Ohne diese Klasse gingen "1., 2., 3." verloren - bei
    Aufgaben fatal, weil sich Lösungen und Materialien darauf beziehen."""

    def __init__(self, dokument):
        self.stufen = {}      # abstractNumId -> {Ebene: (Format, Text, Start)}
        self.num = {}         # numId -> abstractNumId
        self.zaehler = {}     # numId -> {Ebene: Zaehlerstand}
        try:
            wurzel = dokument.part.numbering_part.element
        except Exception:
            return
        for abstrakt in wurzel.findall(W + "abstractNum"):
            stufen = {}
            for ebene in abstrakt.findall(W + "lvl"):
                try:
                    nummer = int(ebene.get(W + "ilvl"))
                    start = int(_wert(ebene.find(W + "start")) or 1)
                except (TypeError, ValueError):
                    continue
                stufen[nummer] = (_wert(ebene.find(W + "numFmt")) or "decimal",
                                  _wert(ebene.find(W + "lvlText")) or "",
                                  start)
            self.stufen[abstrakt.get(W + "abstractNumId")] = stufen
        for num in wurzel.findall(W + "num"):
            verweis = num.find(W + "abstractNumId")
            if verweis is not None:
                self.num[num.get(W + "numId")] = _wert(verweis)

    @staticmethod
    def _numpr(absatz):
        """numId und Ebene - am Absatz selbst oder geerbt aus der Formatvorlage."""
        kandidaten = [absatz._p.pPr]
        stil = absatz.style
        tiefe = 0
        while stil is not None and tiefe < 6:
            kandidaten.append(getattr(stil.element, "pPr", None))
            stil = getattr(stil, "base_style", None)
            tiefe += 1
        num_id, ebene = None, None
        for ppr in kandidaten:
            if ppr is None:
                continue
            numpr = ppr.find(W + "numPr")
            if numpr is None:
                continue
            if num_id is None:
                num_id = _wert(numpr.find(W + "numId"))
            if ebene is None and numpr.find(W + "ilvl") is not None:
                ebene = int(_wert(numpr.find(W + "ilvl")) or 0)
            if num_id is not None:
                break
        return num_id, (ebene or 0)

    def praefix(self, absatz):
        """(Zeichen, Ebene) fuer diesen Absatz oder None. Zeichen ist "-" fuer
        Aufzaehlungspunkte, sonst die ausgerechnete Nummer, z. B. "2." oder "b)"."""
        try:
            num_id, ebene = self._numpr(absatz)
        except Exception:
            return None
        if not num_id or num_id == "0":
            return None
        stufen = self.stufen.get(self.num.get(num_id), {})
        format_, muster, start = stufen.get(ebene, ("decimal", "%%%d." % (ebene + 1), 1))
        zaehler = self.zaehler.setdefault(num_id, {})
        zaehler[ebene] = zaehler.get(ebene, start - 1) + 1
        for tiefer in [k for k in zaehler if k > ebene]:
            del zaehler[tiefer]
        if format_ == "bullet":
            return ("-", ebene)
        if format_ == "none" or not muster:
            return None

        def einsetzen(treffer):
            stufe = int(treffer.group(1)) - 1
            f, _, s = stufen.get(stufe, ("decimal", "", 1))
            return _zahl_formatieren(zaehler.get(stufe, s), f)

        return (re.sub(r"%(\d)", einsetzen, muster).strip(), ebene)


def _innerhalb(element, bis, tag):
    """Liegt element (unterhalb von bis) in einem Element mit diesem Tag?"""
    for vorfahr in element.iterancestors():
        if vorfahr is bis:
            return False
        if vorfahr.tag == tag or vorfahr.tag.endswith("}" + tag.split("}")[-1]):
            return True
    return False


def _laeufe_zu_text(absatz, dokument, sammler, in_ueberschrift, nummerierung):
    """Text eines Absatzes mit fett/kursiv/unterstrichen, Links und Bildern.
    Unterstrichenes bleibt als <u>…</u> erhalten - in Arbeitsblaettern ist
    es oft Teil der Aufgabe ("Ersetzen Sie den unterstrichenen Ausdruck").

    Gibt zurueck: (Text, Bilder als [(Name, Hinweis)], Textfeld-Zeilen)."""
    from docx.text.paragraph import Paragraph

    stuecke = []   # (text, fett, kursiv, unterstrichen) oder ("__BILD__", name, hinweis)
    kaesten = []
    try:
        inhalt = list(absatz.iter_inner_content())
    except AttributeError:
        inhalt = list(absatz.runs)

    for teil in inhalt:
        url = getattr(teil, "url", None)
        if url is not None:                         # Hyperlink
            text = teil.text
            if text.strip():
                stuecke.append(("[%s](%s)" % (text, url), False, False, False))
            continue
        lauf = teil

        # Textfelder. Word speichert sie oft doppelt: als moderne Form und als
        # Rueckfall fuer alte Programme (mc:Fallback) - nur einmal nehmen.
        for feld in lauf._r.iter(W + "txbxContent"):
            if _innerhalb(feld, lauf._r, "Fallback") or _innerhalb(feld, lauf._r, W + "txbxContent"):
                continue
            for p in feld.findall(W + "p"):
                kaesten.extend(_absatz_zu_md(Paragraph(p, dokument), dokument,
                                             sammler, nummerierung))

        for blip in lauf._r.iter(A + "blip"):
            if _innerhalb(blip, lauf._r, W + "txbxContent"):
                continue                            # gehoert zum Textfeld
            rid = blip.get(R + "embed")
            if not rid:
                continue
            name = bx.aus_word_bild(dokument, rid, sammler)
            if not name:
                continue
            hinweis = ""
            zuschnitt = blip.getparent().find(A + "srcRect")
            if zuschnitt is not None and any(
                    (zuschnitt.get(s) or "0") not in ("0", "") for s in ("l", "t", "r", "b")):
                hinweis = ("ACHTUNG: In Word zugeschnitten – das Bild hier ist vollständig "
                           "und zeigt mehr als im Original sichtbar. Seitenansicht prüfen.")
            stuecke.append(("__BILD__", name, hinweis))

        text = lauf.text or ""
        if not text:
            continue
        stuecke.append((
            text,
            bool(lauf.bold) and not in_ueberschrift,
            bool(lauf.italic),
            bool(lauf.underline),
        ))

    # Gleich formatierte Nachbarn zusammenlegen, sonst entsteht **a****b**
    zusammen = []
    for stueck in stuecke:
        if (zusammen and stueck[0] != "__BILD__" and zusammen[-1][0] != "__BILD__"
                and zusammen[-1][1:] == stueck[1:]):
            zusammen[-1] = (zusammen[-1][0] + stueck[0],) + stueck[1:]
        else:
            zusammen.append(stueck)

    ergebnis, bilder = [], []
    for stueck in zusammen:
        if stueck[0] == "__BILD__":
            bilder.append((stueck[1], stueck[2]))
            continue
        text, fett, kursiv, unter = stueck
        kern = text.strip()
        if not kern or SCHREIBLINIE.match(kern):
            ergebnis.append(text)
            continue
        vorne = text[:len(text) - len(text.lstrip())]
        hinten = text[len(text.rstrip()):]
        if unter:
            kern = "<u>%s</u>" % kern
        if kursiv:
            kern = "*%s*" % kern
        if fett:
            kern = "**%s**" % kern
        ergebnis.append(vorne + kern + hinten)
    return "".join(ergebnis), bilder, kaesten


def _absatz_zu_md(absatz, dokument, sammler, nummerierung=None):
    ebene = _stil_ebene(absatz.style)
    text, bilder, kaesten = _laeufe_zu_text(absatz, dokument, sammler, ebene > 0, nummerierung)
    text = text.replace("\t", "  ")
    # Die Nummer wird auch fuer leere Absaetze gezaehlt, sonst verrutscht sie
    praefix = nummerierung.praefix(absatz) if nummerierung else None
    zeilen = []

    if text.strip():
        teile = [t.rstrip() for t in text.split("\n")]
        teile = schreibzeilen_zusammenfassen(teile)
        if ebene > 0:
            kopf = " ".join(t.strip() for t in teile if t.strip())
            if praefix and praefix[0] not in ("-", ""):
                kopf = praefix[0] + " " + kopf       # "1. Einleitung"
            zeilen.append("#" * ebene + " " + kopf)
        elif ebene == -1:
            zeilen.append("*%s*" % " ".join(t.strip() for t in teile if t.strip()))
        else:
            erste = teile[0].lstrip()
            # Aufzaehlungszeichen, auch wenn es fett oder kursiv gesetzt war
            erste = re.sub(r"^(\*\*|\*)?[•▪◦●■](\*\*|\*)?\s*", "- ", erste)
            if praefix and not re.match(r"^(-|\d+[.)]|[a-zA-Z][.)])\s", erste):
                zeichen, tiefe = praefix
                erste = "   " * tiefe + zeichen + " " + erste
            if erste.startswith("#"):
                erste = "\\" + erste
            teile[0] = erste
            # Konvention des Plugins: Ein Zeilenumbruch innerhalb eines
            # Absatzes bleibt ein Zeilenumbruch - so, wie die Lehrkraft es
            # beim Tippen erwartet. md_zu_docx.py setzt das um. Kein
            # Markdown-Doppelleerzeichen: das ist unsichtbar und geht beim
            # Bearbeiten verloren.
            zeilen.append("\n".join(t for t in teile))
    for name, hinweis in bilder:
        zeilen.append(bildzeile(name, hinweis))
    if kaesten:
        # Textfelder als Kasten, wie Einzelzellen-Tabellen
        inhalt = "\n\n".join(kaesten)
        zeilen.append("")                       # Leerzeile vor dem Kasten
        zeilen.append("\n".join("> " + z if z.strip() else ">" for z in inhalt.split("\n")))
    return zeilen


def _tabelle_zu_md(tabelle):
    zeilen = []
    for zeile in tabelle.rows:
        zellen, vorige = [], None
        for zelle in zeile.cells:
            # Verbundene Zellen liefert python-docx mehrfach - nur einmal nehmen
            if vorige is not None and zelle._tc is vorige:
                continue
            vorige = zelle._tc
            zellen.append("\n".join(a.text for a in zelle.paragraphs).strip())
        zeilen.append(zellen)
    if len(zeilen) == 1 and len(zeilen[0]) == 1:
        # Eine einzelne Zelle ist fast immer ein Kasten, keine Tabelle
        inhalt = zeilen[0][0]
        return "\n".join("> " + z if z.strip() else ">" for z in inhalt.split("\n"))
    return pipe_tabelle(zeilen)


def word_umwandeln(pfad, sammler):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    dokument = docx.Document(str(pfad))
    nummerierung = Nummerierung(dokument)
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    bloecke, ueberschriften, stile = [], [], Counter()
    tabellen = 0

    def durchlaufen(element):
        nonlocal tabellen
        for kind in element.iterchildren():
            if kind.tag == w + "p":
                absatz = Paragraph(kind, dokument)
                if absatz.text.strip():
                    stile[absatz.style.name] += 1
                zeilen = _absatz_zu_md(absatz, dokument, sammler, nummerierung)
                if zeilen and zeilen[0].startswith("#"):
                    ueberschriften.append(zeilen[0].lstrip("#").strip())
                bloecke.append("\n".join(zeilen) if zeilen else "")
            elif kind.tag == w + "tbl":
                tabellen += 1
                bloecke.append(_tabelle_zu_md(Table(kind, dokument)))
            elif kind.tag == w + "sdt":
                inhalt = kind.find(w + "sdtContent")
                if inhalt is not None:
                    durchlaufen(inhalt)

    durchlaufen(dokument.element.body)

    # In Word ist jede Schreiblinie ein eigener Absatz - aufeinanderfolgende
    # zu einem einzigen [Schreibzeilen: n] zusammenfassen
    text = "\n\n".join(b for b in bloecke if b.strip())
    text = re.sub(r"\[Schreibzeilen: \d+\](?:\n\n\[Schreibzeilen: \d+\])+",
                  _schreibzeilen_summe, text)

    kopfzeilen = []
    for abschnitt in dokument.sections:
        for bereich in (abschnitt.header, abschnitt.footer):
            try:
                for absatz in bereich.paragraphs:
                    if absatz.text.strip() and absatz.text.strip() not in kopfzeilen:
                        kopfzeilen.append(absatz.text.strip())
            except Exception:
                pass

    return text, {
        "absaetze": sum(stile.values()),
        "tabellen": tabellen,
        "ueberschriften": ueberschriften,
        "formatvorlagen": dict(stile.most_common()),
        "kopf_und_fusszeilen": kopfzeilen,
    }


def _schreibzeilen_summe(treffer):
    anzahl = sum(int(n) for n in re.findall(r"\[Schreibzeilen: (\d+)\]", treffer.group(0)))
    return "[Schreibzeilen: %d]" % anzahl


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

SATZENDE = re.compile(r"[.!?:;…\"“”»)]\s*$")
AUFZAEHLUNG = re.compile(r"^(\d+[.)]|[a-z][.)]|[-•▪◦●■–])\s")


def _schriftmerkmale(seite):
    """Holt je Textstueck Schriftgroesse, -schnitt und Position. Gibt zurueck:
    ein Woerterbuch 'Text -> Merkmale' (zum Einordnen der Zeilen aus
    extract_text()), die Liste der Stuecke und die Liste der Textpositionen
    (fuer die Frage, ob Text ueber einem Bild liegt)."""
    stuecke, punkte = [], []

    def besucher(text, cm, tm, schrift, groesse):
        if not text or not text.strip():
            return
        try:
            punkte.append((cm[0] * tm[4] + cm[2] * tm[5] + cm[4],
                           cm[1] * tm[4] + cm[3] * tm[5] + cm[5],
                           text.strip()))
        except Exception:
            pass
        try:
            skala = math.hypot(tm[0], tm[1]) * math.hypot(cm[0], cm[1])
        except Exception:
            skala = 1
        name = ""
        if schrift:
            try:
                name = str(schrift.get("/BaseFont", ""))
            except Exception:
                name = ""
        for teil in text.split("\n"):
            if teil.strip():
                stuecke.append((teil.strip(), round(groesse * skala, 1),
                                bool(re.search(r"Bold|Black|Heavy|Semibold", name)),
                                bool(re.search(r"Italic|Oblique", name))))

    try:
        seite.extract_text(visitor_text=besucher)
    except Exception:
        return {}, [], []
    merkmale = {}
    for text, groesse, fett, kursiv in stuecke:
        merkmale.setdefault(text, (groesse, fett, kursiv))
    return merkmale, stuecke, punkte


def _matrix_mal(m, n):
    """PDF-Matrizen [a b c d e f] multiplizieren: erst m, dann n."""
    return (m[0] * n[0] + m[1] * n[2], m[0] * n[1] + m[1] * n[3],
            m[2] * n[0] + m[3] * n[2], m[2] * n[1] + m[3] * n[3],
            m[4] * n[0] + m[5] * n[2] + n[4], m[4] * n[1] + m[5] * n[3] + n[5])


def _bildflaechen(seite):
    """Wo auf der Seite liegt welches Bild? {Objektname: [Rechteck, ...]}.
    Verfolgt dazu die Transformationen im Inhaltsstrom der Seite - ein Bild
    wird als Einheitsquadrat gezeichnet, gestreckt durch die aktuelle Matrix."""
    try:
        from pypdf.generic import ContentStream
        inhalt = seite.get_contents()
        if inhalt is None:
            return {}
        befehle = ContentStream(inhalt, seite.pdf).operations
        objekte = seite["/Resources"]["/XObject"]
    except Exception:
        return {}
    matrix, stapel, flaechen = (1, 0, 0, 1, 0, 0), [], {}
    for argumente, befehl in befehle:
        if befehl == b"q":
            stapel.append(matrix)
        elif befehl == b"Q":
            matrix = stapel.pop() if stapel else matrix
        elif befehl == b"cm":
            try:
                matrix = _matrix_mal([float(a) for a in argumente], matrix)
            except Exception:
                pass
        elif befehl == b"Do":
            name = str(argumente[0])
            try:
                if objekte[name].get_object().get("/Subtype") != "/Image":
                    continue
            except Exception:
                continue
            ecken = [(matrix[0] * x + matrix[2] * y + matrix[4],
                      matrix[1] * x + matrix[3] * y + matrix[5])
                     for x, y in ((0, 0), (1, 0), (0, 1), (1, 1))]
            flaechen.setdefault(name, []).append(
                (min(p[0] for p in ecken), min(p[1] for p in ecken),
                 max(p[0] for p in ecken), max(p[1] for p in ecken)))
    return flaechen


def _text_ueber_bild(flaechen, punkte):
    """Text, dessen Position innerhalb eines Bildes liegt. Solcher Text ist
    entweder daraufgelegt (dann fehlt er im herausgeloesten Bild) oder das
    Bild verdeckt ihn - in beiden Faellen zeigt das Bild allein nicht, was
    die Schueler sehen."""
    befund = {}
    for name, rechtecke in flaechen.items():
        treffer = []
        for x, y, text in punkte:
            for x0, y0, x1, y1 in rechtecke:
                if x0 + 1 < x < x1 - 1 and y0 + 1 < y < y1 - 1:
                    treffer.append(text)
                    break
        if treffer:
            text = " ".join(treffer)
            # PDFs setzen Buchstaben manchmal einzeln: "GE O P S ER" -> "GEOPSER".
            # Nur zusammenziehen, wenn die Grossbuchstaben-Stuecke im Schnitt
            # so kurz sind - sonst wuerden echte Woerter verschmolzen.
            gross = [t for t in text.split() if t.isalpha() and t.isupper()]
            if gross and sum(len(t) for t in gross) / len(gross) <= 3:
                text = re.sub(r"(?<=[A-ZÄÖÜ]) (?=[A-ZÄÖÜ])", "", text)
            text = re.sub(r"(?<=_) (?=_)", "", text)
            befund[name] = text
    return befund


def _merkmale_fuer(zeile, merkmale):
    kern = zeile.strip()
    if kern in merkmale:
        return merkmale[kern]
    for text, werte in merkmale.items():
        if kern.startswith(text) or text.startswith(kern):
            return werte
    return None


def _pdf_seite_zu_md(text, merkmale, grundgroesse, ueberschriftgroessen):
    roh = [z.rstrip() for z in text.split("\n")]

    # Umbrochene Schreiblinien ("S________" + "_") wieder zusammensetzen
    zeilen = []
    for zeile in roh:
        if zeile.strip() and set(zeile.strip()) <= {"_"} and zeilen and zeilen[-1].rstrip().endswith("_"):
            zeilen[-1] = zeilen[-1].rstrip() + zeile.strip()
        else:
            zeilen.append(zeile)

    koerper = [z for z in zeilen if len(z.strip()) > 20]
    ueblich = sorted(len(z.strip()) for z in koerper)[len(koerper) // 2] if koerper else 60

    def ebene(nummer, zeile):
        kern = zeile.strip()
        werte = _merkmale_fuer(kern, merkmale)
        if werte and grundgroesse and werte[0] >= grundgroesse * 1.2:
            for rang, groesse in enumerate(ueberschriftgroessen, start=1):
                if abs(werte[0] - groesse) < 0.6:
                    return min(rang, 3)
            return 1
        vorher = zeilen[nummer - 1].strip() if nummer > 0 else ""
        nach_absatzende = not vorher or bool(SATZENDE.search(vorher))
        # Kurze, ganz fette Zeile nach einem Absatzende: Zwischenueberschrift,
        # auch mit Fragezeichen ("Was heißt das jetzt für mich…?"). Nummerierte
        # Aufgaben sind oft ebenfalls fett - die schliesst AUFZAEHLUNG aus.
        if (werte and werte[1] and len(kern) < 80 and nach_absatzende
                and not re.search(r"[.,;:]$", kern)
                and not AUFZAEHLUNG.match(kern) and not SCHREIBLINIE.match(kern)):
            return 2
        # Gleich grosse Zwischenueberschrift: kurz, ohne Satzzeichen am Ende,
        # mindestens zwei Woerter, danach ein langer Absatz
        if (len(kern) < 70 and len(kern.split()) >= 2 and not SATZENDE.search(kern)
                and not AUFZAEHLUNG.match(kern) and not SCHREIBLINIE.match(kern)
                and kern[:1].isupper()):
            vorher = zeilen[nummer - 1].strip() if nummer > 0 else ""
            nachher = zeilen[nummer + 1].strip() if nummer + 1 < len(zeilen) else ""
            if (not vorher or SATZENDE.search(vorher)) and len(nachher) > 0.7 * ueblich:
                return 2
        return 0

    bloecke, absatz, absatz_merkmale = [], [], []

    def absatz_schliessen():
        if not absatz:
            return
        text = absatz[0]
        for zeile in absatz[1:]:
            if re.search(r"[a-zäöüß]-$", text) and zeile[:1].islower():
                text = text[:-1] + zeile            # Silbentrennung aufheben
            else:
                text = text + " " + zeile
        text = re.sub(r"\s{2,}", " ", text).strip()
        schnitte = [m for m in absatz_merkmale if m]
        if schnitte and len(schnitte) == len(absatz_merkmale):
            if all(m[2] for m in schnitte):
                text = "*%s*" % text
            elif all(m[1] for m in schnitte) and len(text) > 70:
                text = "**%s**" % text
        text = re.sub(r"^[•▪◦●■]\s*", "- ", text)
        bloecke.append(text)
        del absatz[:]
        del absatz_merkmale[:]

    for nummer, zeile in enumerate(zeilen):
        kern = zeile.strip()
        if not kern:
            absatz_schliessen()
            continue
        if SCHREIBLINIE.match(kern):
            absatz_schliessen()
            bloecke.append("__LINIE__")
            continue
        stufe = ebene(nummer, zeile)
        if stufe:
            absatz_schliessen()
            bloecke.append("#" * stufe + " " + kern)
            continue
        if absatz:
            vorige = absatz[-1]
            neuer_absatz = (
                AUFZAEHLUNG.match(kern)
                or (SATZENDE.search(vorige) and len(vorige) < 0.8 * ueblich)
            )
            if neuer_absatz:
                absatz_schliessen()
        absatz.append(kern)
        absatz_merkmale.append(_merkmale_fuer(kern, merkmale))
    absatz_schliessen()

    # Linien-Marker zu [Schreibzeilen: n]
    aus, linien = [], 0
    for block in bloecke + [None]:
        if block == "__LINIE__":
            linien += 1
            continue
        if linien:
            aus.append("[Schreibzeilen: %d]" % linien)
            linien = 0
        if block is not None:
            aus.append(block)
    return "\n\n".join(aus)


def pdf_umwandeln(pfad, sammler):
    import pypdf
    leser = pypdf.PdfReader(str(pfad))
    seiten_md, zeichen, scanseiten, ueberschriften = [], 0, [], []

    # Grundschrift und Ueberschriftengroessen ueber das ganze Dokument
    alle_merkmale, alle_stuecke, alle_punkte = [], [], []
    for seite in leser.pages:
        merkmale, stuecke, punkte = _schriftmerkmale(seite)
        alle_merkmale.append(merkmale)
        alle_stuecke.extend(stuecke)
        alle_punkte.append(punkte)
    gewicht = Counter()
    for text, groesse, _, _ in alle_stuecke:
        gewicht[groesse] += len(text)
    grundgroesse = gewicht.most_common(1)[0][0] if gewicht else None
    ueberschriftgroessen = sorted(
        {s[1] for s in alle_stuecke if grundgroesse and s[1] >= grundgroesse * 1.2},
        reverse=True)

    ueberdeckt = {}
    for nummer, seite in enumerate(leser.pages, start=1):
        try:
            text = seite.extract_text() or ""
        except Exception:
            text = ""
        bilder = bx.aus_pdf_seite(seite, nummer, sammler)
        ueber = _text_ueber_bild(_bildflaechen(seite), alle_punkte[nummer - 1])
        zeichen += len(text.strip())
        teile = ["<!-- Seite %d -->" % nummer]
        if len(text.strip()) < 40:
            scanseiten.append(nummer)
            teile.append("[Prüfen: Seite %d ist vermutlich ein Scan. Text aus dem Bild "
                         "abschreiben und diesen Hinweis ersetzen.]" % nummer)
        else:
            md = _pdf_seite_zu_md(text, alle_merkmale[nummer - 1], grundgroesse,
                                  ueberschriftgroessen)
            ueberschriften.extend(z.lstrip("#").strip()
                                  for z in md.split("\n") if z.startswith("#"))
            teile.append(md)
        for objekt, name in bilder:
            hinweis = "Position im Original: Seite %d" % nummer
            if objekt in ueber:
                # Fall 1 aus UMWANDLUNG.md: Das herausgeloeste Bild kann einen
                # alten, verdeckten Stand zeigen. Nur die Seitenansicht zaehlt.
                ueberdeckt[name] = ueber[objekt]
                hinweis += (" · ACHTUNG: Auf der Seite liegt Text über diesem Bild "
                            "(„%s“). Das Bild allein zeigt nicht, was die Schüler "
                            "sehen – Seitenansicht prüfen." % ueber[objekt][:80])
            teile.append(bildzeile(name, hinweis))
        seiten_md.append("\n\n".join(teile))

    return "\n\n".join(seiten_md), {
        "seiten": len(leser.pages),
        "zeichen": zeichen,
        "scanseiten": scanseiten,
        "scan": bool(leser.pages) and len(scanseiten) == len(leser.pages),
        "ueberschriften": ueberschriften,
        "grundschrift_pt": grundgroesse,
        "bilder_mit_text_darueber": ueberdeckt,
    }


# --------------------------------------------------------------------------
# PowerPoint
# --------------------------------------------------------------------------

def pptx_umwandeln(pfad, sammler):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    praesentation = Presentation(str(pfad))
    teile, ueberschriften = [], []

    def rechteck(form):
        return ((form.left or 0), (form.top or 0),
                (form.left or 0) + (form.width or 0), (form.top or 0) + (form.height or 0))

    def ueberlappen(a, b):
        return a[0] < b[2] and b[0] < a[2] and a[1] < b[3] and b[1] < a[3]

    def bildhinweis(form, folie):
        """Folien sind Ebenen: Was spaeter in der Reihenfolge kommt, liegt
        oben. Liegt Text ueber einem Bild oder ist es zugeschnitten, zeigt
        das herausgeloeste Bild nicht, was auf der Folie zu sehen ist."""
        hinweise = []
        try:
            if any((getattr(form, "crop_" + s, 0) or 0) > 0.001
                   for s in ("left", "right", "top", "bottom")):
                hinweise.append("in PowerPoint zugeschnitten – das Bild hier zeigt mehr "
                                "als auf der Folie")
            formen = list(folie.shapes)
            if form in formen:
                bild = rechteck(form)
                darueber = [f for f in formen[formen.index(form) + 1:]
                            if getattr(f, "has_text_frame", False) and f.has_text_frame
                            and f.text_frame.text.strip() and ueberlappen(bild, rechteck(f))]
                if darueber:
                    hinweise.append("auf der Folie liegt Text darüber („%s“)"
                                    % darueber[0].text_frame.text.strip()[:60])
        except Exception:
            pass
        if not hinweise:
            return ""
        return "ACHTUNG: %s. Bild allein nicht verlässlich – Seitenansicht prüfen." \
            % "; ".join(hinweise)

    def formen_durchlaufen(formen, nummer, aus, titelform, folie):
        geordnet = sorted(formen, key=lambda f: ((f.top or 0), (f.left or 0)))
        for form in geordnet:
            if form is titelform:
                continue
            if form.shape_type == MSO_SHAPE_TYPE.GROUP:
                formen_durchlaufen(form.shapes, nummer, aus, titelform, folie)
            elif form.shape_type == MSO_SHAPE_TYPE.PICTURE:
                name = bx.aus_pptx_form(form, nummer, sammler)
                if name:
                    aus.append(bildzeile(name, bildhinweis(form, folie)))
            elif getattr(form, "has_table", False) and form.has_table:
                aus.append(pipe_tabelle([[z.text for z in zeile.cells]
                                         for zeile in form.table.rows]))
            elif getattr(form, "has_text_frame", False) and form.has_text_frame:
                for absatz in form.text_frame.paragraphs:
                    text = "".join(l.text for l in absatz.runs).strip()
                    if text:
                        aus.append("  " * (absatz.level or 0) + "- " + text)

    for nummer, folie in enumerate(praesentation.slides, start=1):
        titelform = folie.shapes.title
        titel = titelform.text_frame.text.strip() if titelform is not None and titelform.has_text_frame else ""
        ueberschriften.append(titel)
        aus = ["## Folie %d%s" % (nummer, (": " + titel) if titel else "")]
        inhalt = []
        formen_durchlaufen(folie.shapes, nummer, inhalt, titelform, folie)
        aus.append("\n".join(inhalt))
        if folie.has_notes_slide:
            notizen = folie.notes_slide.notes_text_frame.text.strip()
            if notizen:
                aus.append("\n".join("> " + z for z in ("Notizen: " + notizen).split("\n")))
        teile.append("\n\n".join(a for a in aus if a.strip()))

    return "\n\n".join(teile), {"folien": len(praesentation.slides),
                                 "ueberschriften": [u for u in ueberschriften if u]}


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def xlsx_umwandeln(pfad, _sammler):
    import openpyxl
    mappe = openpyxl.load_workbook(str(pfad), data_only=True, read_only=True)
    teile, blaetter, gekuerzt = [], [], []
    for blatt in mappe.worksheets:
        zeilen = []
        for zeile in blatt.iter_rows(values_only=True):
            zeilen.append(["" if w is None else str(w) for w in zeile][:30])
            if len(zeilen) >= 200:
                gekuerzt.append(blatt.title)
                break
        while zeilen and not any(c.strip() for c in zeilen[-1]):
            zeilen.pop()
        blaetter.append(blatt.title)
        teile.append("## Tabelle: %s\n\n%s" % (blatt.title, pipe_tabelle(zeilen) or "*(leer)*"))
    if gekuerzt:
        teile.append("[Prüfen: Tabellenblätter nach 200 Zeilen gekürzt: %s]" % ", ".join(gekuerzt))
    return "\n\n".join(teile), {"tabellenblaetter": blaetter, "gekuerzt": gekuerzt}


# --------------------------------------------------------------------------
# Text, Bilder, Umweg ueber LibreOffice
# --------------------------------------------------------------------------

def text_umwandeln(pfad, _sammler):
    roh = Path(pfad).read_bytes()
    for kodierung in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = roh.decode(kodierung)
            break
        except UnicodeDecodeError:
            continue
    vorhandener_kopf, rumpf = g.kopf_trennen(text)
    return rumpf, {"kodierung": kodierung, "vorhandener_kopf": vorhandener_kopf}


def _gerade_richten(pfad, ziel):
    """Handyfotos sind oft quer gespeichert, mit einem Vermerk "beim Anzeigen
    drehen" (EXIF). Viele Programme - und womoeglich Claude - ignorieren den
    Vermerk und sehen das Bild seitlich. Deshalb hier drehen und ohne Vermerk
    speichern. Gibt True zurueck, wenn gedreht wurde."""
    try:
        from PIL import Image, ImageOps
        with Image.open(str(pfad)) as bild:
            if bild.getexif().get(0x0112, 1) in (1, None):
                return False
            gedreht = ImageOps.exif_transpose(bild)
            einstellungen = {"quality": 95} if bild.format == "JPEG" else {}
            gedreht.save(str(ziel), format=bild.format, **einstellungen)
            return True
    except Exception:
        return False


def bild_umwandeln(pfad, sammler):
    daten = Path(pfad).read_bytes()
    # Eigenstaendige Bilder nie wegen Groesse aussortieren - sie SIND der Inhalt
    name = "%s-01%s" % (sammler.praefix, Path(pfad).suffix.lower())
    gedreht = _gerade_richten(pfad, sammler.ziel / name)
    if gedreht:
        daten = (sammler.ziel / name).read_bytes()
    else:
        (sammler.ziel / name).write_bytes(daten)
    masse = bx._masse(daten)
    sammler.gespeichert.append({"datei": name, "herkunft": "eigenständige Bilddatei",
                                "bytes": len(daten),
                                "masse": ("%sx%s" % masse) if masse else None,
                                "ansehbar": Path(pfad).suffix.lower() != ".heic"})
    text = "\n\n".join([
        "![Bild: beschreiben](bilder/%s)" % name,
        "## Beschreibung",
        "[Prüfen: Was zeigt das Bild? Zwei bis vier Sätze.]",
        "## Text im Bild",
        "[Prüfen: Allen lesbaren Text vollständig abschreiben, in der Gliederung "
        "des Bildes. Ohne Text im Bild: diesen Abschnitt löschen.]",
    ])
    return text, {"bild": True, "masse": sammler.gespeichert[-1]["masse"],
                  "gerade_gerichtet": gedreht}


def libreoffice_finden():
    for name in ("soffice", "libreoffice"):
        pfad = shutil.which(name)
        if pfad:
            return pfad
    for ort in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                "/usr/bin/soffice"):
        if Path(ort).exists():
            return ort
    return None


def umweg_libreoffice(pfad, zielformat):
    """Alte und fremde Formate (.doc, .ppt, .odt ...) erst in das moderne
    Office-Format bringen. LibreOffice ist in beiden Cowork-Umgebungen
    vorhanden (Befund T1)."""
    soffice = libreoffice_finden()
    if not soffice:
        raise RuntimeError("LibreOffice fehlt – %s-Dateien können nicht gelesen werden"
                           % Path(pfad).suffix)
    zwischen = Path(tempfile.mkdtemp(prefix="ua-umweg-"))
    subprocess.run([soffice, "--headless", "--norestore", "--convert-to", zielformat,
                    "--outdir", str(zwischen), str(pfad)],
                   stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=300)
    ergebnis = zwischen / (Path(pfad).stem + "." + zielformat)
    if not ergebnis.exists():
        raise RuntimeError("LibreOffice konnte %s nicht umwandeln" % Path(pfad).name)
    return ergebnis


# --------------------------------------------------------------------------
# Seitenansicht und Vollstaendigkeit
# --------------------------------------------------------------------------

def ansicht_erzeugen(lesbar, endung, bereich):
    """Legt ansicht.pdf in den Arbeitsbereich: das Dokument so, wie die
    Schueler es sehen. Claude prueft daran die Abschrift; das Skript
    vergleicht damit die Vollstaendigkeit.

    Word wird bewusst NICHT ueber das PDF umgewandelt - die Word-Datei kennt
    Ueberschriften, Tabellen und Unterstreichungen, das PDF nur Buchstaben an
    Positionen (UMWANDLUNG.md Kap. 3). Das PDF dient nur zum Vergleich."""
    ziel = Path(bereich) / "ansicht.pdf"
    if endung == ".pdf":
        shutil.copy2(str(lesbar), str(ziel))
        return {"datei": ziel.name, "weg": "Original"}
    if endung not in (".docx", ".pptx"):
        return {"datei": None, "grund": "für diese Dateiart nicht nötig"}

    soffice = libreoffice_finden()
    if soffice:
        zwischen = Path(tempfile.mkdtemp(prefix="ua-ansicht-"))
        subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "pdf",
                        "--outdir", str(zwischen), str(lesbar)],
                       stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=300)
        erzeugt = zwischen / (Path(lesbar).stem + ".pdf")
        if erzeugt.exists():
            shutil.move(str(erzeugt), str(ziel))
            return {"datei": ziel.name, "weg": "LibreOffice"}
    if endung == ".docx":
        # Rueckfall auf installiertes MS Word - nur ausserhalb der Sandbox
        try:
            from docx2pdf import convert
            convert(str(lesbar), str(ziel))
            if ziel.exists():
                return {"datei": ziel.name, "weg": "MS Word"}
        except Exception:
            pass
    return {"datei": None, "fehler": "Keine Seitenansicht möglich – LibreOffice fehlt"}


WORT = re.compile(r"[^\W\d_]{2,}")
NUMMER_AM_ANFANG = re.compile(r"^[\s*_>#\-]*(\d{1,2}|[a-z])[.)]\s")


def _nummern_zaehlen(text):
    """Zaehlt Nummerierungen ("1.", "b)") am Anfang von Zeilen und
    Tabellenzellen. Formatierung davor (**, >, -) zaehlt nicht mit - sonst
    gaelte eine fett gesetzte Aufgabe "**1. …**" als unnummeriert."""
    return sum(1 for stueck in re.split(r"\n|\||<br>", text) if NUMMER_AM_ANFANG.match(stueck))


def _woerter(text):
    text = unicodedata.normalize("NFKC", text)      # Ligaturen: ﬁ -> fi
    return Counter(w.lower() for w in WORT.findall(text))


def vollstaendigkeit(ansicht, rumpf, ausnahmen=()):
    """Vergleicht die Woerter der Seitenansicht mit denen der Abschrift.

    Faengt Luecken, die der Word- oder PowerPoint-Leser hat, ohne dass jede
    einzeln bekannt sein muss: Textfelder, automatische Nummern, Fussnoten,
    SmartArt. Ist die Abschrift auffaellig unvollstaendig, bekommt sie oben
    einen Pruefhinweis - der sperrt das Ablegen, bis Claude nachgesehen hat."""
    import pypdf
    leser = pypdf.PdfReader(str(ansicht))
    text = "\n".join((s.extract_text() or "") for s in leser.pages)
    text = unicodedata.normalize("NFKC", text)
    # Silbentrennung am Zeilenende aufheben, sonst fehlen zerteilte Woerter
    text = re.sub(r"([a-zäöüß])-\s*\n\s*([a-zäöüß])", r"\1\2", text)

    fehlend = _woerter(text) - _woerter(rumpf) - _woerter(" ".join(ausnahmen))
    anzahl = sum(fehlend.values())
    nummern_ansicht = _nummern_zaehlen(text)
    nummern_abschrift = _nummern_zaehlen(rumpf)
    nummern_fehlen = max(0, nummern_ansicht - nummern_abschrift)

    auffaellig = anzahl >= 3 or nummern_fehlen >= 2
    teile = []
    if anzahl:
        teile.append("Diese Wörter stehen in der Seitenansicht, aber nicht hier: %s"
                     % ", ".join(w for w, _ in fehlend.most_common(30)))
    if nummern_fehlen >= 2:
        teile.append("Nummerierungen am Zeilenanfang: in der Ansicht %d, hier %d"
                     % (nummern_ansicht, nummern_abschrift))
    hinweis = ("[Prüfen: In der Abschrift fehlen vermutlich Teile des Originals "
               "(Seitenansicht: ansicht.pdf). %s. Fehlendes ergänzen, dann diesen "
               "Hinweis löschen.]" % ". ".join(teile)) if auffaellig else None
    return {
        "fehlende_woerter": anzahl,
        "beispiele": [w for w, _ in fehlend.most_common(30)],
        "nummern_ansicht": nummern_ansicht,
        "nummern_abschrift": nummern_abschrift,
        "auffaellig": auffaellig,
        "hinweis": hinweis,
    }


# --------------------------------------------------------------------------
# Eine Datei
# --------------------------------------------------------------------------

def umwandeln(quelle, arbeitsbereich):
    """Wandelt eine Datei in arbeitsbereich/inhalt.md um. Gibt den Bericht
    zurueck; wirft keine Ausnahme nach aussen."""
    quelle = Path(quelle)
    arbeitsbereich = Path(arbeitsbereich)
    (arbeitsbereich / "bilder").mkdir(parents=True, exist_ok=True)
    sammler = bx.Sammler(arbeitsbereich / "bilder", g.kurzname(quelle.stem, 30))
    endung = quelle.suffix.lower()
    bericht = {"quelle": quelle.name, "endung": endung}

    try:
        lesbar = quelle
        if endung in UMWEG_WORD:
            lesbar, endung = umweg_libreoffice(quelle, "docx"), ".docx"
        elif endung in UMWEG_PPTX:
            lesbar, endung = umweg_libreoffice(quelle, "pptx"), ".pptx"
        elif endung in UMWEG_XLSX:
            lesbar, endung = umweg_libreoffice(quelle, "xlsx"), ".xlsx"
        if lesbar is not quelle:
            bericht["umweg"] = "über LibreOffice in %s umgewandelt" % endung

        wandler = {
            ".docx": ("word", word_umwandeln),
            ".pdf": ("pdf", pdf_umwandeln),
            ".pptx": ("powerpoint", pptx_umwandeln),
            ".xlsx": ("excel", xlsx_umwandeln), ".xlsm": ("excel", xlsx_umwandeln),
            ".txt": ("text", text_umwandeln), ".md": ("text", text_umwandeln),
        }
        if endung in wandler:
            art, funktion = wandler[endung]
        elif endung in BILDDATEIEN:
            art, funktion = "bild", bild_umwandeln
        else:
            raise RuntimeError("Dateiart %s wird nicht unterstützt" % endung)

        rumpf, einzelheiten = funktion(lesbar, sammler)
        bericht["art"] = art
        bericht.update(einzelheiten)
    except Exception as fehler:
        bericht["fehler"] = "{}: {}".format(type(fehler).__name__, fehler)
        return bericht

    bericht["bilder"] = sammler.bericht()
    bericht["zeichen_md"] = len(rumpf)
    loesung = [u for u in bericht.get("ueberschriften", []) if LOESUNGSTEIL.search(u)]
    if loesung:
        bericht["loesungsteil_ab"] = loesung[0]

    # Seitenansicht und Vollstaendigkeit (UMWANDLUNG.md Kap. 6, Fall 1).
    # Fehler hier duerfen die Aufnahme nicht verhindern - sie fehlen dann nur.
    try:
        bericht["ansicht"] = ansicht_erzeugen(lesbar, endung, arbeitsbereich)
    except Exception as fehler:
        bericht["ansicht"] = {"datei": None, "fehler": str(fehler)}
    if bericht["ansicht"].get("datei"):
        try:
            pruefung = vollstaendigkeit(arbeitsbereich / bericht["ansicht"]["datei"], rumpf,
                                        bericht.get("kopf_und_fusszeilen", []))
            bericht["vollstaendigkeit"] = pruefung
            if pruefung["auffaellig"]:
                rumpf = pruefung["hinweis"] + "\n\n" + rumpf
        except Exception as fehler:
            bericht["vollstaendigkeit"] = {"fehler": str(fehler)}

    felder = {
        "schema": g.SCHEMA, "fach": None, "klasse": None, "thema": None,
        "typ": None, "beschreibung": None, "herkunft": None, "quelle": None,
        "original": "originale/" + quelle.name,
        "status": "pruefen",
        "grafiken": "ja" if sammler.gespeichert else "nein",
        "erstellt": g.heute(), "aktualisiert": g.heute(),
    }
    # Hatte eine mitgebrachte .md schon einen Kopf, dessen Werte uebernehmen
    for schluessel, wert in (bericht.get("vorhandener_kopf") or {}).items():
        if wert not in (None, "") and schluessel not in ("schema", "original", "erstellt"):
            felder[schluessel] = wert

    g.md_schreiben(arbeitsbereich / "inhalt.md", felder, aufraeumen(rumpf))
    return bericht


# --------------------------------------------------------------------------
# Stapel aus eingang/
# --------------------------------------------------------------------------

def eingang_dateien(wurzel):
    eingang = wurzel / "eingang"
    if not eingang.is_dir():
        return []
    dateien = []
    for pfad in sorted(eingang.rglob("*")):
        if not pfad.is_file():
            continue
        name = pfad.name
        if (name.lower() in IGNORIEREN or name.startswith((".", "~$"))
                or any(t.startswith(".") for t in pfad.relative_to(eingang).parts)):
            continue
        dateien.append(pfad)
    return dateien


def vorhandene_auftraege(wurzel):
    """Pruefwerte der Dateien, die schon in Bearbeitung sind."""
    bereits = {}
    for auftrag in (wurzel / "_system" / "aufnahme").glob("*/auftrag.json"):
        try:
            daten = json.loads(auftrag.read_text(encoding="utf-8"))
            bereits[daten["pruefwert"]] = auftrag.parent.name
        except Exception:
            continue
    return bereits


def stapel(wurzel, anzahl):
    wurzel = Path(wurzel).resolve()
    aufnahme = wurzel / "_system" / "aufnahme"
    aufnahme.mkdir(parents=True, exist_ok=True)
    bereits = vorhandene_auftraege(wurzel)
    # Schon aufgenommene Originale gar nicht erst umwandeln - das spart Claude
    # die Arbeit und der Lehrkraft Kontingent. ablegen.py prueft zur
    # Sicherheit noch einmal.
    aufgenommen = {e.get("pruefwert"): md
                   for md, e in g.stand_lesen(wurzel).get("aufgenommen", {}).items()}

    ergebnisse, neu = [], 0
    offen = eingang_dateien(wurzel)
    for quelle in offen:
        pruef = g.pruefwert(quelle)
        if pruef in aufgenommen:
            ergebnisse.append({"quelle": g.relativ(quelle, wurzel),
                               "schon_aufgenommen": aufgenommen[pruef]})
            continue
        if pruef in bereits:
            ergebnisse.append({"quelle": g.relativ(quelle, wurzel),
                               "arbeitsbereich": "_system/aufnahme/" + bereits[pruef],
                               "schon_in_arbeit": True})
            continue
        if neu >= anzahl:
            continue
        name = g.kurzname(quelle.stem, 50)
        bereich = aufnahme / name
        nummer = 2
        while bereich.exists():
            bereich = aufnahme / ("%s-%d" % (name, nummer))
            nummer += 1
        bericht = umwandeln(quelle, bereich)
        auftrag = {
            "quelle": g.relativ(quelle, wurzel),
            "pruefwert": pruef,
            "umgewandelt": g.jetzt(),
            "bericht": bericht,
        }
        (bereich / "auftrag.json").write_text(
            json.dumps(auftrag, indent=2, ensure_ascii=False), encoding="utf-8")
        eintrag = {"quelle": auftrag["quelle"],
                   "arbeitsbereich": g.relativ(bereich, wurzel)}
        eintrag.update(bericht)
        ergebnisse.append(eintrag)
        neu += 1

    in_arbeit = [e for e in ergebnisse
                 if not e.get("schon_in_arbeit") and not e.get("schon_aufgenommen")]
    return {
        "arbeitsordner": str(wurzel),
        "portion": anzahl,
        "umgewandelt": len([e for e in in_arbeit if not e.get("fehler")]),
        "fehlgeschlagen": len([e for e in in_arbeit if e.get("fehler")]),
        "schon_in_arbeit": len([e for e in ergebnisse if e.get("schon_in_arbeit")]),
        "schon_aufgenommen": len([e for e in ergebnisse if e.get("schon_aufgenommen")]),
        "warten_noch": max(0, len(offen) - len(ergebnisse)),
        "dateien": ergebnisse,
    }


def text_bericht(daten):
    z = ["Umwandlung aus dem Eingang", "=" * 60]
    for e in daten["dateien"]:
        if e.get("schon_in_arbeit"):
            z.append("  (schon in Arbeit)  %s" % e["quelle"])
            continue
        if e.get("schon_aufgenommen"):
            z.append("  DOPPELT  %s – ist schon aufgenommen als %s"
                     % (e["quelle"], e["schon_aufgenommen"]))
            continue
        if e.get("fehler"):
            z.append("  FEHLER  %s – %s" % (e["quelle"], e["fehler"]))
            continue
        merkmale = []
        if e.get("seiten"):
            merkmale.append("%d Seiten" % e["seiten"])
        if e.get("folien"):
            merkmale.append("%d Folien" % e["folien"])
        if e.get("tabellen"):
            merkmale.append("%d Tabellen" % e["tabellen"])
        bilder = len((e.get("bilder") or {}).get("gespeichert", []))
        if bilder and not e.get("bild"):
            merkmale.append("1 Bild" if bilder == 1 else "%d Bilder" % bilder)
        if e.get("scan"):
            merkmale.append("SCAN – Text muss aus dem Bild gelesen werden")
        elif e.get("scanseiten"):
            merkmale.append("Scan-Seiten: %s" % e["scanseiten"])
        if e.get("bild"):
            merkmale.append("Bilddatei – ansehen und beschreiben")
        if e.get("loesungsteil_ab"):
            merkmale.append("Lösungsteil ab „%s“" % e["loesungsteil_ab"])
        if e.get("bilder_mit_text_darueber"):
            anzahl = len(e["bilder_mit_text_darueber"])
            merkmale.append("Text über Bild: %d – Seitenansicht maßgeblich" % anzahl)
        if (e.get("vollstaendigkeit") or {}).get("auffaellig"):
            merkmale.append("UNVOLLSTÄNDIG? %d Wörter der Ansicht fehlen"
                            % e["vollstaendigkeit"]["fehlende_woerter"])
        if e.get("gerade_gerichtet"):
            merkmale.append("Foto gerade gerichtet")
        ansicht = e.get("ansicht") or {}
        if ansicht.get("fehler"):
            merkmale.append("keine Seitenansicht (%s)" % ansicht["fehler"])
        z.append("  ok  %s  (%s)" % (e["quelle"], ", ".join(merkmale) or e.get("art")))
        z.append("      → %s/inhalt.md" % e["arbeitsbereich"])
    z.append("")
    z.append("Umgewandelt: %d · Fehler: %d · warten noch im Eingang: %d"
             % (daten["umgewandelt"], daten["fehlgeschlagen"], daten["warten_noch"]))
    return "\n".join(z)


def main():
    parser = argparse.ArgumentParser(description="Dateien aus eingang/ in Markdown umwandeln")
    gruppe = parser.add_mutually_exclusive_group(required=True)
    gruppe.add_argument("--ordner", help="Arbeitsordner der Lehrkraft (Stapelbetrieb)")
    gruppe.add_argument("--datei", help="Eine einzelne Datei umwandeln (zum Testen)")
    parser.add_argument("--ziel", help="Zielordner für --datei")
    parser.add_argument("--anzahl", type=int, default=STANDARD_PORTION,
                        help="Höchstens so viele neue Dateien pro Durchlauf (Standard 10)")
    parser.add_argument("--json", action="store_true")
    argumente = parser.parse_args()

    if argumente.datei:
        ziel = Path(argumente.ziel or (Path(argumente.datei).stem + "-md"))
        bericht = umwandeln(argumente.datei, ziel)
        print(json.dumps(bericht, indent=2, ensure_ascii=False, default=str))
        return 0

    daten = stapel(argumente.ordner, argumente.anzahl)
    g.ausgeben(daten, argumente.json, text_bericht)
    return 0


if __name__ == "__main__":
    try:
        ENDE = main()
    except BaseException:
        import traceback
        traceback.print_exc()
        ENDE = 1
    sys.exit(ENDE)
