#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Laufzeit-Test fuer den Unterrichtsassistenten (Phase 1a).

Beantwortet die Testpunkte T1, T2, T3 und T6 aus KONZEPT.md Kap. 13 mit
harten Fakten: Es erzeugt echte Dateien und liest sie zurueck. Eine
erzeugte Datei beweist nichts - nur das Zuruecklesen.

T4, T5, T7, T8 und T9 kann kein Skript beantworten. Die traegt der Skill
'laufzeit-test' als Beobachtungen in denselben Bericht nach.

Laeuft mit reiner Standardbibliothek: Das Skript muss starten, BEVOR
klar ist, ob ueberhaupt etwas installiert ist. Kein Probenfehler bricht
den Lauf ab, jeder Fehlschlag ist selbst ein Ergebnis. Exit-Code immer 0.

Aufruf:
    python laufzeit_test.py --umgebung "Claude Cowork"
    python laufzeit_test.py --umgebung "Claude Cowork" --install
"""

import argparse
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import traceback
import urllib.error
import urllib.request
from datetime import datetime
from pathlib import Path

try:  # Umlaute und Ampel-Symbole auch auf Windows-Konsolen
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

SCHEMA = 1

# (Importname, pip-Name, wofuer der Unterrichtsassistent es braucht, unverzichtbar?)
BIBLIOTHEKEN = [
    ("docx",       "python-docx",    "Word erzeugen und lesen (md_zu_docx.py)",        True),
    ("pptx",       "python-pptx",    "PowerPoint erzeugen (md_zu_pptx.py)",            True),
    ("pypdf",      "pypdf",          "PDF lesen, Seiten zaehlen",                      True),
    ("PIL",        "Pillow",         "Bilder pruefen und zuschneiden",                 True),
    ("yaml",       "PyYAML",         "YAML-Metadatenkopf lesen (Kap. 5.3)",            True),
    ("markdown",   "Markdown",       "Markdown zu HTML (Ersatzweg PDF, HTML-Folien)",  True),
    ("fitz",       "PyMuPDF",        "Text und Bilder aus PDF (zu_markdown.py)",       False),
    ("openpyxl",   "openpyxl",       "Excel lesen (zu_markdown.py)",                   False),
    ("xhtml2pdf",  "xhtml2pdf",      "HTML zu PDF, braucht keine Systemprogramme",     False),
    ("fpdf",       "fpdf2",          "PDF direkt bauen, braucht keine Systemprogramme", False),
    ("reportlab",  "reportlab",      "PDF direkt bauen, braucht keine Systemprogramme", False),
    ("weasyprint", "weasyprint",     "HTML zu PDF, braucht GTK-Systembibliotheken",    False),
    ("docx2pdf",   "docx2pdf",       "Word zu PDF ueber installiertes Office",         False),
    ("lxml",       "lxml",           "XML-Unterbau von python-docx/-pptx",             False),
    ("jinja2",     "Jinja2",         "HTML-Folien aus Vorlage (md_zu_folien.py)",      False),
    ("bs4",        "beautifulsoup4", "HTML auswerten (Rechtsstand-Pruefung)",          False),
    ("requests",   "requests",       "Webzugriff bequemer als urllib",                 False),
    ("pandas",     "pandas",         "Tabellen (optional)",                            False),
]

# Pakete, die nur fuer den Weg zum PDF gebraucht werden. Sie gelten oben als
# optional, werden mit --install aber trotzdem probiert: ohne sie laesst sich
# T2 nicht ehrlich beantworten.
#
# Reihenfolge nach Anspruch an die Umgebung. Die ersten drei sind reine
# pip-Pakete und brauchen weder LibreOffice noch Word, LaTeX oder GTK - sie
# muessen laufen, sonst haengt die PDF-Ausgabe an einer Fremdinstallation auf
# dem Rechner der Lehrkraft. Die letzten beiden duerfen scheitern: weasyprint
# braucht Systembibliotheken, docx2pdf ein installiertes Office.
PDF_PAKETE = ["xhtml2pdf", "fpdf2", "reportlab", "weasyprint", "docx2pdf"]

# Externe Programme, die fuer die Umwandlung in Frage kommen
WERKZEUGE = ["soffice", "libreoffice", "pandoc", "wkhtmltopdf", "pdflatex", "git", "pip"]

# Orte, an denen LibreOffice ueblicherweise liegt, ohne im PATH zu stehen
LIBREOFFICE_ORTE = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
    "/usr/local/bin/soffice",
    "/snap/bin/libreoffice",
]

WEB_ZIELE = [
    ("https://www.gesetze-im-internet.de/bgb/__622.html", "622", "Gesetzestext BGB (Kap. 8.2)"),
    ("https://www.gesetze-im-internet.de/kschg/", "ndigungsschutzgesetz", "Gesetzesuebersicht"),
]

# Bewusst mit Umlauten und Paragraphzeichen: Genau die gehen auf dem Weg
# .md -> HTML -> PDF gern verloren, und ein Arbeitsblatt fuer Arbeit & Recht
# besteht zur Haelfte daraus. Ein Beispiel in ASCII wuerde nichts beweisen.
BEISPIEL_MD = """---
schema: 1
fach: arbeit-recht
klasse: "11"
thema: Kündigungsschutz
typ: arbeitsblatt
---

# Arbeitsblatt: Kündigungsfristen

## Material M1: Sachverhalt

Frau König ist seit vier Monaten bei der Firma X beschäftigt. Sie
erhält eine Kündigung mit einer Frist von zwei Wochen.

## Aufgabe 1

Prüfen Sie anhand von M1, ob die Kündigung fristgerecht ist
(§ 622 Abs. 3 BGB).

| Kriterium | Punkte |
| --- | --- |
| Frist in der Probezeit erkannt | 2 |
| Zugang bestimmt | 2 |
"""


# --------------------------------------------------------------------------
# Hilfsmittel
# --------------------------------------------------------------------------

def jetzt():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def versuch(fn, *args, **kwargs):
    """Fuehrt fn aus und faengt wirklich alles ab. Nie darf eine Probe den Lauf beenden."""
    try:
        return {"ok": True, "wert": fn(*args, **kwargs), "fehler": None}
    except BaseException as fehler:  # auch SystemExit aus fremden Bibliotheken
        return {
            "ok": False,
            "wert": None,
            "fehler": "{}: {}".format(type(fehler).__name__, fehler),
            "spur": traceback.format_exc(limit=3).strip().splitlines()[-3:],
        }


def lauf(befehl, timeout=180, cwd=None):
    """Externes Programm aufrufen, Ausgabe gekuerzt zurueckgeben."""
    start = datetime.now()
    fertig = subprocess.run(
        befehl,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        timeout=timeout,
        cwd=str(cwd) if cwd else None,
    )
    ausgabe = fertig.stdout.decode("utf-8", errors="replace").strip()
    return {
        "befehl": " ".join(str(teil) for teil in befehl),
        "code": fertig.returncode,
        "sekunden": round((datetime.now() - start).total_seconds(), 1),
        "ausgabe": ausgabe[-1500:],
    }


def pdf_seiten(pfad):
    """Seitenzahl eines PDF. Erst sauber per pypdf, sonst grob ueber die Rohdaten."""
    pfad = Path(pfad)
    if not pfad.exists():
        return None
    try:
        import pypdf
        return len(pypdf.PdfReader(str(pfad)).pages)
    except Exception:
        pass
    try:
        daten = pfad.read_bytes()
        treffer = re.findall(rb"/Type\s*/Page[^s]", daten)
        return len(treffer) or None
    except Exception:
        return None


def pdf_gueltig(pfad):
    """Faengt die Datei mit %PDF an und hat sie ueberhaupt Inhalt?"""
    pfad = Path(pfad)
    return pfad.exists() and pfad.stat().st_size > 400 and pfad.read_bytes()[:5] == b"%PDF-"


def pdf_text(pfad):
    """Steht der Text im PDF als Text - oder ist er zu Pixeln geworden?
    Fuer ein Arbeitsblatt entscheidend: sonst laesst es sich nicht durchsuchen,
    nicht kopieren und nicht vorlesen."""
    try:
        import pypdf
        inhalt = "\n".join(
            (seite.extract_text() or "") for seite in pypdf.PdfReader(str(pfad)).pages)
        return {"umlaute": "Kündigung" in inhalt, "paragraphzeichen": "§ 622" in inhalt}
    except Exception as fehler:
        return {"pruefung_nicht_moeglich": "{}: {}".format(type(fehler).__name__, fehler)}


def pdf_ergebnis(ziel, protokoll=None):
    """Einheitlicher Befund fuer jeden PDF-Weg."""
    daten = {
        "datei": datei_info(ziel),
        "gueltig": pdf_gueltig(ziel),
        "seiten": pdf_seiten(ziel),
        "text": pdf_text(ziel) if pdf_gueltig(ziel) else None,
    }
    if protokoll is not None:
        daten["protokoll"] = protokoll
    return daten


def datei_info(pfad):
    pfad = Path(pfad)
    if not pfad.exists():
        return {"vorhanden": False}
    return {"vorhanden": True, "bytes": pfad.stat().st_size, "pfad": str(pfad)}


# --------------------------------------------------------------------------
# Probe E: Umgebung
# --------------------------------------------------------------------------

def probe_umgebung(ordner):
    ergebnis = {
        "python": sys.version.split()[0],
        "python_vollstaendig": " ".join(sys.version.split()),
        "ausfuehrbar": sys.executable,
        "plattform": platform.platform(),
        "system": platform.system(),
        "architektur": platform.machine(),
        "benutzer_ordner": str(Path.home()),
        "arbeitsverzeichnis": str(Path.cwd()),
        "ausgabeordner": str(ordner),
        "temp": tempfile.gettempdir(),
        "pfadtrenner": os.sep,
        "dateikodierung": sys.getfilesystemencoding(),
    }

    # Schreibproben - auch mit Umlauten, denn Materialien heissen "Kündigungsschutz"
    proben = {}
    for name, dateiname in (
        ("ascii", "schreibprobe.txt"),
        ("umlaute", "Kündigungsschutz_Prüfung_Übung.txt"),
    ):
        def schreiben(dateiname=dateiname):
            ziel = ordner / dateiname
            ziel.write_text("Ä Ö Ü ß – § 622 BGB", encoding="utf-8")
            zurueck = ziel.read_text(encoding="utf-8")
            if "§ 622" not in zurueck:
                raise ValueError("Inhalt kam veraendert zurueck: %r" % zurueck)
            return {"name": ziel.name, "bytes": ziel.stat().st_size}
        proben[name] = versuch(schreiben)
    ergebnis["schreibproben"] = proben

    # Lange Pfade - unter Windows historisch bei 260 Zeichen Schluss
    def langer_pfad():
        tief = ordner / ("ordner_" + "x" * 40) / ("ordner_" + "y" * 40) / ("ordner_" + "z" * 40)
        tief.mkdir(parents=True, exist_ok=True)
        ziel = tief / ("datei_" + "w" * 60 + ".txt")
        ziel.write_text("ok", encoding="utf-8")
        return {"laenge": len(str(ziel)), "pfad": str(ziel)}
    ergebnis["langer_pfad"] = versuch(langer_pfad)

    # Externe Programme
    gefunden = {}
    for werkzeug in WERKZEUGE:
        gefunden[werkzeug] = shutil.which(werkzeug)
    for ort in LIBREOFFICE_ORTE:
        if Path(ort).exists() and not gefunden.get("soffice"):
            gefunden["soffice"] = ort
    ergebnis["werkzeuge"] = gefunden

    # Darf das Skript ueberhaupt andere Programme starten?
    ergebnis["unterprozess_erlaubt"] = versuch(
        lauf, [sys.executable, "-c", "print('unterprozess laeuft')"], timeout=60
    )
    return ergebnis


# --------------------------------------------------------------------------
# Probe A: T1 - Bibliotheken
# --------------------------------------------------------------------------

def _version(importname, pipname):
    try:
        from importlib import metadata
        return metadata.version(pipname)
    except Exception:
        pass
    try:
        modul = __import__(importname)
        return str(getattr(modul, "__version__", "unbekannt"))
    except Exception:
        return "unbekannt"


def probe_bibliotheken(nachinstallieren):
    def pruefen():
        stand = {}
        for importname, pipname, zweck, wichtig in BIBLIOTHEKEN:
            try:
                __import__(importname)
                stand[pipname] = {
                    "importname": importname,
                    "zweck": zweck,
                    "unverzichtbar": wichtig,
                    "vorhanden": True,
                    "version": _version(importname, pipname),
                }
            except BaseException as fehler:
                stand[pipname] = {
                    "importname": importname,
                    "zweck": zweck,
                    "unverzichtbar": wichtig,
                    "vorhanden": False,
                    "version": None,
                    "fehler": "{}: {}".format(type(fehler).__name__, fehler),
                }
        return stand

    ergebnis = {"vorher": pruefen(), "installation": None,
                "installation_pdf": None, "nachher": None}

    fehlend_wichtig = [
        pip for pip, wert in ergebnis["vorher"].items()
        if wert["unverzichtbar"] and not wert["vorhanden"]
    ]
    fehlend_pdf = [
        pip for pip in PDF_PAKETE
        if pip in ergebnis["vorher"] and not ergebnis["vorher"][pip]["vorhanden"]
    ]
    ergebnis["fehlend_unverzichtbar"] = fehlend_wichtig
    ergebnis["fehlend_pdf"] = fehlend_pdf

    def installieren(pakete):
        if not pakete:
            return {"ok": True, "wert": None, "hinweis": "nichts zu tun, war schon vorhanden"}
        return versuch(
            lauf,
            [sys.executable, "-m", "pip", "install", "--disable-pip-version-check"] + pakete,
            timeout=900,
        )

    if nachinstallieren:
        ergebnis["installation"] = installieren(fehlend_wichtig)
        # Die Wege zum PDF einzeln, nicht als Gruppe: pip installiert entweder
        # alles oder nichts. Ein Paket, das sich nicht aufloesen laesst, wuerde
        # sonst die anderen mitreissen - und damit einen Weg verdecken, der
        # funktioniert haette.
        ergebnis["installation_pdf"] = {
            paket: installieren([paket]) for paket in fehlend_pdf
        }
        # Importcache leeren, sonst sieht Python die frische Installation nicht
        namen = {eintrag[0] for eintrag in BIBLIOTHEKEN}
        for modul in list(sys.modules):
            if modul.split(".")[0] in namen:
                sys.modules.pop(modul, None)
        ergebnis["nachher"] = pruefen()

    return ergebnis


# --------------------------------------------------------------------------
# Probe B: T3 - Word und PowerPoint mit Vorlage
# --------------------------------------------------------------------------

def probe_word(ordner):
    """Vorlage mit eigenen Formatvorlagen bauen, daraus ein Dokument erzeugen und
    zurueckgelesen pruefen, ob die Formatvorlagen wirklich angewendet wurden."""
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Pt

    vorlage = ordner / "vorlage.docx"
    ergebnisdatei = ordner / "arbeitsblatt.docx"
    schritte = {}

    # 1. Vorlage mit benannten Formatvorlagen - so wie die Schule sie liefern wuerde
    dokument = docx.Document()
    for name, groesse, fett in (("UA-Titel", 18, True), ("UA-Aufgabe", 12, True), ("UA-Text", 11, False)):
        stil = dokument.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        stil.font.size = Pt(groesse)
        stil.font.bold = fett
    dokument.save(str(vorlage))
    schritte["vorlage_erzeugt"] = datei_info(vorlage)

    # 2. Vorlage oeffnen und befuellen - genau das macht md_zu_docx.py spaeter
    aus_vorlage = docx.Document(str(vorlage))
    aus_vorlage.add_paragraph("Arbeitsblatt: Kündigungsfristen", style="UA-Titel")
    aus_vorlage.add_paragraph("Aufgabe 1 [Punkte: 6]", style="UA-Aufgabe")
    aus_vorlage.add_paragraph(
        "Prüfen Sie anhand von M1, ob die Kündigung fristgerecht ist. § 622 BGB", style="UA-Text")
    tabelle = aus_vorlage.add_table(rows=2, cols=2)
    tabelle.cell(0, 0).text = "Kriterium"
    tabelle.cell(0, 1).text = "Punkte"
    tabelle.cell(1, 0).text = "Frist erkannt"
    tabelle.cell(1, 1).text = "2"
    aus_vorlage.add_page_break()
    aus_vorlage.add_paragraph("Lösungsversion", style="UA-Titel")
    aus_vorlage.save(str(ergebnisdatei))
    schritte["dokument_erzeugt"] = datei_info(ergebnisdatei)

    # 3. Zurueckgelesen - erst das beweist etwas
    kontrolle = docx.Document(str(ergebnisdatei))
    stile = [absatz.style.name for absatz in kontrolle.paragraphs if absatz.text.strip()]
    texte = [absatz.text for absatz in kontrolle.paragraphs if absatz.text.strip()]
    schritte["rueckgelesen"] = {
        "absaetze": len(texte),
        "verwendete_formatvorlagen": sorted(set(stile)),
        "eigene_formatvorlagen_angewendet": sorted({s for s in stile if s.startswith("UA-")}),
        "tabellen": len(kontrolle.tables),
        "tabelle_inhalt_erhalten": bool(kontrolle.tables)
            and kontrolle.tables[0].cell(1, 0).text == "Frist erkannt",
        "umlaute_erhalten": any("Kündigung" in t for t in texte),
        "paragraphzeichen_erhalten": any("§ 622" in t for t in texte),
    }

    schritte["bewertung"] = (
        len(schritte["rueckgelesen"]["eigene_formatvorlagen_angewendet"]) == 3
        and schritte["rueckgelesen"]["tabelle_inhalt_erhalten"]
        and schritte["rueckgelesen"]["umlaute_erhalten"]
        and schritte["rueckgelesen"]["paragraphzeichen_erhalten"]
    )
    schritte["datei"] = str(ergebnisdatei)
    return schritte


def probe_powerpoint(ordner):
    """Gleiche Logik: Vorlage erzeugen, als Vorlage wiederverwenden, zurueckpruefen."""
    from pptx import Presentation
    from pptx.util import Pt

    vorlage = ordner / "vorlage.pptx"
    ergebnisdatei = ordner / "praesentation.pptx"
    schritte = {}

    grundlage = Presentation()
    schritte["layouts_der_vorlage"] = [layout.name for layout in grundlage.slide_layouts]
    grundlage.save(str(vorlage))
    schritte["vorlage_erzeugt"] = datei_info(vorlage)

    aus_vorlage = Presentation(str(vorlage))
    layout = aus_vorlage.slide_layouts[1]  # Titel und Inhalt
    folie = aus_vorlage.slides.add_slide(layout)
    folie.shapes.title.text = "Kündigungsschutz – Überblick"
    rahmen = folie.placeholders[1].text_frame
    rahmen.text = "Probezeit: zwei Wochen Frist (§ 622 Abs. 3 BGB)"
    zweite = rahmen.add_paragraph()
    zweite.text = "Danach: gesetzliche Grundfrist"
    zweite.font.size = Pt(18)
    aus_vorlage.save(str(ergebnisdatei))
    schritte["datei_erzeugt"] = datei_info(ergebnisdatei)

    kontrolle = Presentation(str(ergebnisdatei))
    texte = []
    for folie in kontrolle.slides:
        for form in folie.shapes:
            if form.has_text_frame:
                texte.append(form.text_frame.text)
    schritte["rueckgelesen"] = {
        "folien": len(kontrolle.slides),
        "texte": texte,
        "umlaute_erhalten": any("Kündigung" in t for t in texte),
        "paragraphzeichen_erhalten": any("§ 622" in t for t in texte),
    }
    schritte["bewertung"] = (
        len(kontrolle.slides) == 1
        and schritte["rueckgelesen"]["umlaute_erhalten"]
        and schritte["rueckgelesen"]["paragraphzeichen_erhalten"]
    )
    schritte["datei"] = str(ergebnisdatei)
    return schritte


# --------------------------------------------------------------------------
# Probe C: T2 - PDF auf beiden Wegen
# --------------------------------------------------------------------------

def markdown_zu_html(quelle):
    """Nutzt die Bibliothek 'markdown', wenn vorhanden. Sonst ein Minimal-Ersatz,
    denn fuer den Test zaehlt der Weg zum PDF, nicht die Qualitaet der Umwandlung."""
    try:
        import markdown as md
        rumpf = md.markdown(quelle, extensions=["tables"])
        weg = "markdown-Bibliothek"
    except Exception:
        zeilen = []
        for zeile in quelle.splitlines():
            blank = zeile.strip()
            if not blank or blank == "---":
                continue
            if blank.startswith("## "):
                zeilen.append("<h2>%s</h2>" % blank[3:])
            elif blank.startswith("# "):
                zeilen.append("<h1>%s</h1>" % blank[2:])
            elif blank.startswith("|"):
                zeilen.append("<p><code>%s</code></p>" % blank)
            else:
                zeilen.append("<p>%s</p>" % blank)
        rumpf = "\n".join(zeilen)
        weg = "Minimal-Ersatz (Standardbibliothek)"
    html = (
        "<!doctype html>\n<html lang='de'><head><meta charset='utf-8'>"
        "<title>Arbeitsblatt</title>"
        "<style>body{font-family:sans-serif;margin:2cm;} h1{font-size:18pt;}</style>"
        "</head><body>\n%s\n</body></html>\n" % rumpf
    )
    return html, weg


def probe_pdf(ordner, docx_pfad):
    # Drei Kategorien, absteigend nach dem, was die Lehrkraft davon hat:
    #   word_zu_pdf - PDF sieht aus wie die Word-Vorlage der Schule
    #   ersatzweg   - PDF aus HTML, eigenes Layout, aber ohne Fremdprogramme moeglich
    #   direkt      - PDF im Skript gebaut, laeuft immer, Layout kostet Arbeit
    ergebnis = {"word_zu_pdf": {}, "ersatzweg": {}, "direkt": {}}
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        for ort in LIBREOFFICE_ORTE:
            if Path(ort).exists():
                soffice = ort
                break

    # --- Weg 1: Word -> PDF (der Hauptweg aus Kap. 5.2) -------------------
    if docx_pfad and Path(docx_pfad).exists():
        ziel_ordner = ordner / "pdf_aus_word"
        ziel_ordner.mkdir(exist_ok=True)

        def ueber_docx2pdf():
            from docx2pdf import convert
            ziel = ziel_ordner / "docx2pdf.pdf"
            convert(str(docx_pfad), str(ziel))
            return pdf_ergebnis(ziel)
        ergebnis["word_zu_pdf"]["docx2pdf (MS Word)"] = versuch(ueber_docx2pdf)

        if soffice:
            def ueber_libreoffice():
                protokoll = lauf(
                    [soffice, "--headless", "--norestore", "--convert-to", "pdf",
                     "--outdir", str(ziel_ordner), str(docx_pfad)],
                    timeout=300,
                )
                return pdf_ergebnis(ziel_ordner / (Path(docx_pfad).stem + ".pdf"), protokoll)
            ergebnis["word_zu_pdf"]["LibreOffice"] = versuch(ueber_libreoffice)
        else:
            ergebnis["word_zu_pdf"]["LibreOffice"] = {
                "ok": False, "wert": None, "fehler": "soffice/libreoffice nicht gefunden"}

        if shutil.which("pandoc"):
            def ueber_pandoc():
                ziel = ziel_ordner / "pandoc.pdf"
                protokoll = lauf(["pandoc", str(docx_pfad), "-o", str(ziel)], timeout=300)
                return pdf_ergebnis(ziel, protokoll)
            ergebnis["word_zu_pdf"]["pandoc"] = versuch(ueber_pandoc)
        else:
            ergebnis["word_zu_pdf"]["pandoc"] = {
                "ok": False, "wert": None, "fehler": "pandoc nicht gefunden"}
    else:
        ergebnis["word_zu_pdf"]["entfaellt"] = {
            "ok": False, "wert": None,
            "fehler": "Keine Word-Datei vorhanden - Probe T3 ist vorher fehlgeschlagen"}

    # --- Weg 2: Ersatzweg .md -> HTML -> PDF (Kap. 5.2) -------------------
    html, umwandlungsweg = markdown_zu_html(BEISPIEL_MD)
    html_datei = ordner / "arbeitsblatt.html"
    html_datei.write_text(html, encoding="utf-8")
    ergebnis["ersatzweg"]["markdown_zu_html"] = {
        "weg": umwandlungsweg, "datei": datei_info(html_datei)}

    ziel_ordner = ordner / "pdf_aus_html"
    ziel_ordner.mkdir(exist_ok=True)

    # xhtml2pdf zuerst: reines pip-Paket, braucht kein einziges Systemprogramm.
    def ueber_xhtml2pdf():
        from xhtml2pdf import pisa
        ziel = ziel_ordner / "xhtml2pdf.pdf"
        with open(str(ziel), "wb") as datei:
            status = pisa.CreatePDF(html, dest=datei, encoding="utf-8")
        if status.err:
            raise RuntimeError("xhtml2pdf meldet %s Fehler" % status.err)
        return pdf_ergebnis(ziel)
    ergebnis["ersatzweg"]["xhtml2pdf (ohne Fremdprogramm)"] = versuch(ueber_xhtml2pdf)

    def ueber_weasyprint():
        from weasyprint import HTML
        ziel = ziel_ordner / "weasyprint.pdf"
        HTML(string=html).write_pdf(str(ziel))
        return pdf_ergebnis(ziel)
    ergebnis["ersatzweg"]["weasyprint"] = versuch(ueber_weasyprint)

    if shutil.which("wkhtmltopdf"):
        def ueber_wkhtmltopdf():
            ziel = ziel_ordner / "wkhtmltopdf.pdf"
            protokoll = lauf(["wkhtmltopdf", str(html_datei), str(ziel)], timeout=300)
            return pdf_ergebnis(ziel, protokoll)
        ergebnis["ersatzweg"]["wkhtmltopdf"] = versuch(ueber_wkhtmltopdf)
    else:
        ergebnis["ersatzweg"]["wkhtmltopdf"] = {
            "ok": False, "wert": None, "fehler": "wkhtmltopdf nicht gefunden"}

    if soffice:
        def html_ueber_libreoffice():
            protokoll = lauf(
                [soffice, "--headless", "--norestore", "--convert-to", "pdf",
                 "--outdir", str(ziel_ordner), str(html_datei)],
                timeout=300,
            )
            return pdf_ergebnis(ziel_ordner / (html_datei.stem + ".pdf"), protokoll)
        ergebnis["ersatzweg"]["LibreOffice (HTML)"] = versuch(html_ueber_libreoffice)

    # --- Weg 3: PDF direkt im Skript bauen -------------------------------
    # Letzte Rueckfallebene. Laeuft ohne jedes Fremdprogramm, aber das Layout
    # muss md_zu_pdf.py dann selbst setzen - kein HTML, keine Word-Vorlage.
    ziel_ordner = ordner / "pdf_direkt"
    ziel_ordner.mkdir(exist_ok=True)
    TITEL_ZEILE = "Arbeitsblatt: Kündigungsfristen"
    TEXT_ZEILE = ("Prüfen Sie anhand von M1, ob die Kündigung fristgerecht ist "
                  "(§ 622 Abs. 3 BGB).")

    def ueber_fpdf2():
        from fpdf import FPDF
        ziel = ziel_ordner / "fpdf2.pdf"
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("helvetica", "B", 18)
        pdf.cell(0, 12, TITEL_ZEILE, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", size=11)
        pdf.multi_cell(0, 7, TEXT_ZEILE)
        pdf.output(str(ziel))
        return pdf_ergebnis(ziel)
    ergebnis["direkt"]["fpdf2 (ohne Fremdprogramm)"] = versuch(ueber_fpdf2)

    def ueber_reportlab():
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Table
        ziel = ziel_ordner / "reportlab.pdf"
        stile = getSampleStyleSheet()
        SimpleDocTemplate(str(ziel), pagesize=A4).build([
            Paragraph(TITEL_ZEILE, stile["Heading1"]),
            Paragraph(TEXT_ZEILE, stile["BodyText"]),
            Table([["Kriterium", "Punkte"], ["Frist in der Probezeit erkannt", "2"]]),
        ])
        return pdf_ergebnis(ziel)
    ergebnis["direkt"]["reportlab (ohne Fremdprogramm)"] = versuch(ueber_reportlab)

    return ergebnis


# --------------------------------------------------------------------------
# Probe D: T6 - Webzugriff
# --------------------------------------------------------------------------

def probe_web():
    ergebnis = {}
    for adresse, marker, zweck in WEB_ZIELE:
        def abrufen(adresse=adresse, marker=marker):
            anfrage = urllib.request.Request(
                adresse, headers={"User-Agent": "unterrichtsassistent-laufzeittest/0.1"})
            with urllib.request.urlopen(anfrage, timeout=20) as antwort:
                rohdaten = antwort.read(200000)
                text = rohdaten.decode("utf-8", errors="replace")
                return {
                    "status": getattr(antwort, "status", None) or antwort.getcode(),
                    "bytes": len(rohdaten),
                    "inhalt_plausibel": marker in text,
                }
        ergebnis[adresse] = versuch(abrufen)
        ergebnis[adresse]["zweck"] = zweck
    return ergebnis


# --------------------------------------------------------------------------
# Bewertung: Ampel je Testpunkt
# --------------------------------------------------------------------------

def bewerten(daten):
    ampel = {}

    # T1 Bibliotheken
    bib = daten["bibliotheken"]
    stand = bib.get("nachher") or bib["vorher"]
    unverzichtbar = [w for w in stand.values() if w["unverzichtbar"]]
    da = [w for w in unverzichtbar if w["vorhanden"]]
    if unverzichtbar and len(da) == len(unverzichtbar):
        nachtrag = " (nach Nachinstallation)" if bib.get("nachher") else ""
        ampel["T1"] = ("ok", "Alle %d unverzichtbaren Bibliotheken verfügbar%s" % (len(da), nachtrag))
    elif da:
        fehlt = ", ".join(sorted(k for k, w in stand.items() if w["unverzichtbar"] and not w["vorhanden"]))
        ampel["T1"] = ("teilweise", "%d von %d verfügbar, es fehlen: %s" % (len(da), len(unverzichtbar), fehlt))
    else:
        ampel["T1"] = ("fehlt", "Keine der unverzichtbaren Bibliotheken verfügbar")

    # T2 PDF
    pdf = daten.get("pdf", {})

    def geglueckt(bereich):
        treffer = []
        for name, wert in pdf.get(bereich, {}).items():
            if isinstance(wert, dict) and wert.get("ok") and isinstance(wert.get("wert"), dict):
                if wert["wert"].get("gueltig"):
                    treffer.append(name)
        return treffer

    word_wege = geglueckt("word_zu_pdf")
    ersatz_wege = geglueckt("ersatzweg")
    direkt_wege = geglueckt("direkt")
    if word_wege:
        ampel["T2"] = ("ok", "Word → PDF funktioniert über: %s. Das PDF folgt der "
                             "Word-Vorlage der Schule." % ", ".join(word_wege))
    elif ersatz_wege:
        ampel["T2"] = ("teilweise",
                       "Word → PDF scheitert. Ersatzweg .md → HTML → PDF geht über: %s. "
                       "Das PDF sieht dann anders aus als die Word-Datei."
                       % ", ".join(ersatz_wege))
    elif direkt_wege:
        ampel["T2"] = ("teilweise",
                       "Nur der direkte Weg geht (%s). PDFs sind möglich, aber das "
                       "Layout muss ein Skript selbst setzen." % ", ".join(direkt_wege))
    else:
        ampel["T2"] = ("fehlt", "Kein Weg zu einem PDF gefunden – auch keiner, "
                                "der ohne Fremdprogramme auskommt")

    # T3 Word/PowerPoint mit Vorlage
    word = daten.get("word", {})
    pptx = daten.get("powerpoint", {})
    word_ok = bool(word.get("ok")) and bool((word.get("wert") or {}).get("bewertung"))
    pptx_ok = bool(pptx.get("ok")) and bool((pptx.get("wert") or {}).get("bewertung"))
    if word_ok and pptx_ok:
        ampel["T3"] = ("ok", "Word und PowerPoint aus Vorlage erzeugt, Rücklesen bestätigt")
    elif word_ok or pptx_ok:
        ampel["T3"] = ("teilweise", "Nur %s funktioniert" % ("Word" if word_ok else "PowerPoint"))
    else:
        ampel["T3"] = ("fehlt", "Weder Word noch PowerPoint konnten erzeugt werden")

    # T6 Web
    web = daten.get("web", {})
    erreicht = [a for a, w in web.items() if w.get("ok") and (w.get("wert") or {}).get("inhalt_plausibel")]
    if web and len(erreicht) == len(web):
        ampel["T6"] = ("ok", "Alle %d offiziellen Quellen im Skript erreichbar" % len(web))
    elif erreicht:
        ampel["T6"] = ("teilweise", "%d von %d Quellen erreichbar" % (len(erreicht), len(web)))
    else:
        ampel["T6"] = ("fehlt", "Kein Webzugriff aus dem Skript heraus")

    for punkt in ("T4", "T5", "T7", "T8", "T9"):
        ampel[punkt] = ("offen", "Beobachtung – trägt der Skill nach")
    return ampel


SYMBOL = {"ok": "✅", "teilweise": "⚠️", "fehlt": "❌", "offen": "⬜"}

TITEL = {
    "T1": "Welche Python-Bibliotheken stehen zur Verfügung bzw. lassen sich installieren?",
    "T2": "Funktioniert Word → PDF? Sonst Ersatzweg .md → HTML → PDF?",
    "T3": "Funktioniert die Erzeugung von PowerPoint und Word mit Vorlage?",
    "T4": "Wie heißen die Befehle tatsächlich? Wie werden Skripte aus Skills aufgerufen?",
    "T5": "Kann ein Skill oder Sub-Agent ein Modell festlegen?",
    "T6": "Hat Cowork Webzugriff (für die Rechtsstand-Prüfung)?",
    "T7": "Kann das Plugin die Ordner-Anweisung automatisch setzen?",
    "T8": "Greift der passende Skill bei freien Formulierungen ohne Befehl?",
    "T9": "Wie viele Dateien pro /aufnehmen-Durchlauf sind sinnvoll?",
    "T6b": "Darf Claude selbst im Web suchen (nicht nur das Skript)?",
}

# Fragen, die kein Skript beantworten kann - der Skill traegt sie nach.
BEOBACHTUNGEN = [
    ("T4", [
        "Wie lautet der Befehl, mit dem dieser Skill aufgerufen wurde (genauer Text)?",
        "Gibt es ein Plugin-Präfix, z. B. `/unterrichtsassistent:laufzeit-test`?",
        "Wie wurde das Skript aufgerufen – welcher Pfad hat funktioniert?",
        "Hat `${CLAUDE_PLUGIN_ROOT}` funktioniert?",
        "Musste die Ausführung des Skripts bestätigt werden?",
    ]),
    ("T5", [
        "Mit welchem Modell läuft diese Sitzung?",
        "Ließ sich im Skill ein Modell festlegen (z. B. Feld `model:` im Frontmatter)?",
        "Gibt es in Cowork eine sichtbare Modellauswahl, und wo?",
    ]),
    ("T7", [
        "Konnte das Plugin die Ordner-Anweisung selbst setzen oder eine Datei dafür anlegen?",
        "Falls nein: Wie viele Schritte muss die Lehrkraft von Hand tun?",
    ]),
    ("T8", [
        "Welche freien Formulierungen wurden ausprobiert (bitte wörtlich)?",
        "Bei welchen griff der Skill ohne Befehl, bei welchen nicht?",
    ]),
    ("T9", [
        "Wie lange dauerte dieser Lauf, und wie viel Kontingent hat er verbraucht?",
        "Empfohlene Zahl Dateien pro `/aufnehmen`-Durchlauf (Startwert im Konzept: 10)?",
    ]),
    ("T6b", [
        "Darf Claude selbst in Cowork im Web suchen (nicht nur das Skript)?",
        "Falls ja: Wie heißt das Werkzeug dafür?",
    ]),
]


# --------------------------------------------------------------------------
# Bericht
# --------------------------------------------------------------------------

def bericht_schreiben(daten, ziel):
    ampel = daten["ampel"]
    z = []
    z.append("# Laufzeit-Test Unterrichtsassistent (Phase 1a)")
    z.append("")
    z.append("| | |")
    z.append("| --- | --- |")
    z.append("| **Umgebung** | %s |" % daten["umgebung"])
    z.append("| **Zeitpunkt** | %s |" % daten["zeitpunkt"])
    z.append("| **Python** | %s |" % daten["system"]["python"])
    z.append("| **Plattform** | %s |" % daten["system"]["plattform"])
    z.append("")
    z.append("> Dieser Bericht gilt **nur für die oben genannte Umgebung**. "
             "Ein Lauf in Claude Code sagt nichts über Cowork aus.")
    z.append("")

    z.append("## Ergebnis auf einen Blick")
    z.append("")
    z.append("| Punkt | Frage | Ergebnis |")
    z.append("| --- | --- | --- |")
    for punkt in ("T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8", "T9"):
        stufe, text = ampel[punkt]
        z.append("| **%s** | %s | %s %s |" % (punkt, TITEL[punkt], SYMBOL[stufe], text))
    z.append("")
    z.append("⬜ = kann ein Skript nicht beantworten. Diese Punkte trägt der Skill "
             "`laufzeit-test` unten als Beobachtung nach.")
    z.append("")

    # --- T1 ---
    z.append("## T1 – Bibliotheken")
    z.append("")
    bib = daten["bibliotheken"]
    stand = bib.get("nachher") or bib["vorher"]
    z.append("| Bibliothek | Gebraucht für | Nötig | Status |")
    z.append("| --- | --- | --- | --- |")
    for pipname, wert in stand.items():
        status = "✅ %s" % wert["version"] if wert["vorhanden"] else "❌ fehlt"
        z.append("| `%s` | %s | %s | %s |" % (
            pipname, wert["zweck"], "ja" if wert["unverzichtbar"] else "optional", status))
    z.append("")

    if bib.get("installation"):
        inst = bib["installation"]
        z.append("**Nachinstallation (unverzichtbar):** %s" % (
            inst.get("hinweis") or ("geglückt" if inst.get("ok") else "fehlgeschlagen")))
        if bib.get("fehlend_unverzichtbar"):
            z.append("")
            z.append("Versucht wurde: %s"
                     % ", ".join("`%s`" % p for p in bib["fehlend_unverzichtbar"]))
        if inst.get("fehler"):
            z.append("")
            z.append("Fehler: `%s`" % inst["fehler"])
        z.append("")

        z.append("**Nachinstallation (Wege zum PDF), einzeln versucht:**")
        z.append("")
        einzeln = bib.get("installation_pdf") or {}
        if not einzeln:
            z.append("Nichts zu tun, alle waren schon vorhanden.")
        else:
            z.append("| Paket | Installation |")
            z.append("| --- | --- |")
            for paket, wert in einzeln.items():
                z.append("| `%s` | %s |" % (
                    paket,
                    "✅ geglückt" if wert.get("ok") else "❌ %s" % wert.get("fehler")))
            z.append("")
            z.append("> Installiert heißt noch nicht lauffähig – ob ein Weg wirklich "
                     "ein PDF erzeugt, steht bei T2.")
            z.append(">")
            z.append("> **Dieser Lauf unterschätzt T2.** Nach einer Nachinstallation "
                     "einmal **erneut ohne `--install`** laufen lassen: Manche Pakete "
                     "sind erst in einem frischen Prozess voll nutzbar. `docx2pdf` "
                     "etwa zieht `pywin32` nach und meldet im selben Lauf noch einen "
                     "Fehler, im nächsten erzeugt es das PDF.")
        z.append("")
    else:
        z.append("_Nachinstallation nicht versucht. Für die zweite Hälfte von T1 "
                 "(„bzw. lassen sich installieren“) das Skript erneut mit `--install` aufrufen._")
        z.append("")

    z.append("**Externe Programme:**")
    z.append("")
    z.append("| Programm | Gefunden unter |")
    z.append("| --- | --- |")
    for name, ort in daten["system"]["werkzeuge"].items():
        z.append("| `%s` | %s |" % (name, ("`%s`" % ort) if ort else "– nicht vorhanden"))
    z.append("")

    # --- T3 ---
    z.append("## T3 – Word und PowerPoint aus Vorlage")
    z.append("")
    z.append("Geprüft wird nicht, ob eine Datei entsteht, sondern ob sie **zurückgelesen** "
             "noch stimmt: eigene Formatvorlagen angewendet, Tabelle erhalten, Umlaute und § intakt.")
    z.append("")
    for bezeichnung, schluessel in (("Word", "word"), ("PowerPoint", "powerpoint")):
        probe = daten.get(schluessel, {})
        z.append("### %s" % bezeichnung)
        z.append("")
        if not probe.get("ok"):
            z.append("❌ Fehlgeschlagen: `%s`" % probe.get("fehler"))
            z.append("")
            continue
        wert = probe["wert"]
        z.append("%s Datei: `%s`" % ("✅" if wert["bewertung"] else "⚠️", wert["datei"]))
        z.append("")
        for name, inhalt in wert["rueckgelesen"].items():
            z.append("- **%s:** %s" % (name.replace("_", " "), inhalt))
        z.append("")

    # --- T2 ---
    z.append("## T2 – PDF")
    z.append("")
    z.append("Drei Wege, absteigend nach dem, was die Lehrkraft davon hat. Nur der "
             "Hauptweg liefert ein PDF, das der Word-Vorlage der Schule folgt – "
             "er braucht dafür aber ein Fremdprogramm auf dem Rechner.")
    z.append("")
    for bereich, ueberschrift in (
        ("word_zu_pdf", "Hauptweg: Word → PDF (braucht Word oder LibreOffice)"),
        ("ersatzweg", "Ersatzweg: .md → HTML → PDF"),
        ("direkt", "Rückfallebene: PDF direkt im Skript bauen"),
    ):
        z.append("### %s" % ueberschrift)
        z.append("")
        z.append("| Weg | Ergebnis | Seiten | Text im PDF | Größe |")
        z.append("| --- | --- | --- | --- | --- |")
        for name, wert in daten.get("pdf", {}).get(bereich, {}).items():
            if not isinstance(wert, dict):
                continue
            if name == "markdown_zu_html":
                z.append("| Markdown → HTML | ✅ über %s | – | – | %s Bytes |" % (
                    wert.get("weg"), wert.get("datei", {}).get("bytes", "?")))
                continue
            if wert.get("ok") and isinstance(wert.get("wert"), dict):
                inhalt = wert["wert"]
                text = inhalt.get("text") or {}
                if "pruefung_nicht_moeglich" in text:
                    textspalte = "? nicht prüfbar"
                elif text.get("umlaute") and text.get("paragraphzeichen"):
                    textspalte = "✅ Umlaute und §"
                elif text:
                    textspalte = "⚠️ unvollständig"
                else:
                    textspalte = "–"
                z.append("| %s | %s | %s | %s | %s Bytes |" % (
                    name,
                    "✅ gültiges PDF" if inhalt.get("gueltig")
                    else "⚠️ Datei entstand, ist aber kein gültiges PDF",
                    inhalt.get("seiten") or "?",
                    textspalte,
                    inhalt.get("datei", {}).get("bytes", "?")))
            else:
                z.append("| %s | ❌ %s | – | – | – |" % (
                    name, wert.get("fehler", "fehlgeschlagen")))
        z.append("")
    z.append("„Text im PDF“ prüft, ob der Inhalt als **Text** im PDF steht und nicht "
             "als Pixel – sonst lässt sich ein Arbeitsblatt weder durchsuchen noch "
             "kopieren noch vorlesen.")
    z.append("")

    # --- T6 ---
    z.append("## T6 – Webzugriff aus dem Skript")
    z.append("")
    z.append("| Quelle | Ergebnis |")
    z.append("| --- | --- |")
    for adresse, wert in daten.get("web", {}).items():
        if wert.get("ok"):
            inhalt = wert["wert"]
            z.append("| %s | ✅ Status %s, %s Bytes, Inhalt plausibel: %s |" % (
                adresse, inhalt["status"], inhalt["bytes"],
                "ja" if inhalt["inhalt_plausibel"] else "nein"))
        else:
            z.append("| %s | ❌ %s |" % (adresse, wert.get("fehler")))
    z.append("")
    z.append("> Das prüft nur das Netz **im Skript**. Ob Claude selbst in Cowork "
             "im Web suchen darf, ist eine andere Frage – siehe Beobachtung T6b.")
    z.append("")

    # --- Umgebung ---
    z.append("## Umgebung im Detail")
    z.append("")
    system = daten["system"]
    for name in ("python_vollstaendig", "ausfuehrbar", "plattform", "architektur",
                 "arbeitsverzeichnis", "ausgabeordner", "temp", "dateikodierung"):
        z.append("- **%s:** `%s`" % (name.replace("_", " "), system.get(name)))
    schreib = system.get("schreibproben", {}).get("umlaute", {})
    z.append("- **Schreiben mit Umlauten im Dateinamen:** %s" % (
        "✅ geht" if schreib.get("ok") else "❌ %s" % schreib.get("fehler")))
    langer = system.get("langer_pfad", {})
    z.append("- **Langer Pfad (%s Zeichen):** %s" % (
        (langer.get("wert") or {}).get("laenge", "?"),
        "✅ geht" if langer.get("ok") else "❌ %s" % langer.get("fehler")))
    z.append("- **Andere Programme starten:** %s" % (
        "✅ erlaubt" if system.get("unterprozess_erlaubt", {}).get("ok") else "❌ blockiert"))
    z.append("")

    # --- Beobachtungen ---
    z.append("---")
    z.append("")
    z.append("## Beobachtungen aus der Sitzung")
    z.append("")
    z.append("_Diese Punkte kann kein Skript beantworten. Der Skill `laufzeit-test` "
             "füllt sie aus, während er in Cowork läuft. Solange hier „offen“ steht, "
             "ist Phase 1a **nicht** abgeschlossen._")
    z.append("")
    for punkt, fragen in BEOBACHTUNGEN:
        z.append("### %s – %s" % (punkt, TITEL.get(punkt, "")))
        z.append("")
        for frage in fragen:
            z.append("- **%s**" % frage)
            z.append("  - offen")
        z.append("")

    z.append("---")
    z.append("")
    z.append("## Fazit")
    z.append("")
    z.append("_Trägt der Skill nach: Kann so weitergebaut werden, oder müssen "
             "Architektur und Designprinzip 3 neu bewertet werden (Kap. 11)?_")
    z.append("")

    ziel.write_text("\n".join(z), encoding="utf-8")
    return ziel


# --------------------------------------------------------------------------
# Hauptprogramm
# --------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Laufzeit-Test des Unterrichtsassistenten (Phase 1a, Testpunkte T1-T3, T6)")
    parser.add_argument("--ausgabe", default="laufzeit-test",
                        help="Ordner für Bericht und Testdateien (Standard: ./laufzeit-test)")
    parser.add_argument("--umgebung", default="unbekannt",
                        help='Wo läuft das? z. B. "Claude Cowork" oder "Claude Code (Entwicklung)"')
    parser.add_argument("--install", action="store_true",
                        help="Fehlende unverzichtbare Bibliotheken nachinstallieren (verändert die Umgebung)")
    argumente = parser.parse_args()

    ordner = Path(argumente.ausgabe).expanduser()
    ordner.mkdir(parents=True, exist_ok=True)
    # Erst anlegen, dann aufloesen: resolve() liefert unter Python < 3.10 auf
    # Windows einen relativen Pfad zurueck, wenn der Ordner noch nicht existiert.
    ordner = ordner.resolve()

    print("Laufzeit-Test Unterrichtsassistent – Phase 1a")
    print("Umgebung laut Aufruf: %s" % argumente.umgebung)
    print("Ausgabe: %s" % ordner)
    print("-" * 60)

    daten = {
        "schema": SCHEMA,
        "zeitpunkt": jetzt(),
        "umgebung": argumente.umgebung,
        "aufruf": " ".join(sys.argv),
    }

    print("[1/5] Umgebung prüfen …")
    daten["system"] = probe_umgebung(ordner)

    print("[2/5] T1: Bibliotheken prüfen%s …" % (" und nachinstallieren" if argumente.install else ""))
    daten["bibliotheken"] = probe_bibliotheken(argumente.install)

    print("[3/5] T3: Word und PowerPoint aus Vorlage erzeugen …")
    daten["word"] = versuch(probe_word, ordner)
    daten["powerpoint"] = versuch(probe_powerpoint, ordner)

    print("[4/5] T2: PDF-Wege prüfen (kann dauern, LibreOffice startet langsam) …")
    docx_pfad = daten["word"]["wert"].get("datei") if daten["word"].get("ok") else None
    pdf_probe = versuch(probe_pdf, ordner, docx_pfad)
    daten["pdf"] = pdf_probe["wert"] if pdf_probe["ok"] else {
        "word_zu_pdf": {}, "ersatzweg": {}, "fehler": pdf_probe["fehler"]}

    print("[5/5] T6: Webzugriff prüfen …")
    daten["web"] = probe_web()

    daten["ampel"] = bewerten(daten)

    roh = ordner / "roh.json"
    roh.write_text(json.dumps(daten, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    bericht = bericht_schreiben(daten, ordner / "bericht.md")

    print("-" * 60)
    for punkt in ("T1", "T2", "T3", "T6"):
        stufe, text = daten["ampel"][punkt]
        print("%s %s  %s" % (SYMBOL[stufe], punkt, text))
    print("⬜ T4, T5, T7, T8, T9 – Beobachtungen, trägt der Skill nach")
    print("-" * 60)
    nachinstalliert = daten["bibliotheken"].get("installation_pdf") or {}
    if argumente.install and nachinstalliert:
        print("Hinweis: Es wurde nachinstalliert (%d Pakete für PDF). Dieser Lauf"
              % len(nachinstalliert))
        print("unterschätzt T2 möglicherweise – bitte einmal ERNEUT ohne --install")
        print("starten, damit frisch installierte Pakete voll zur Verfügung stehen.")
        print("-" * 60)
    print("Bericht:  %s" % bericht)
    print("Rohdaten: %s" % roh)
    return 0


if __name__ == "__main__":
    try:
        ENDE = main()
    except BaseException:
        # Ein Absturz ist selbst ein Ergebnis und wird gemeldet. Das sys.exit()
        # steht bewusst ausserhalb: sonst faengt der Block seinen eigenen
        # SystemExit und meldet einen Absturz, der keiner ist.
        traceback.print_exc()
        print("\nDas Testskript selbst ist abgestürzt. Das ist schon ein Befund: "
              "die Umgebung trägt nicht einmal die Standardbibliothek wie erwartet.")
        ENDE = 0
    sys.exit(ENDE)
