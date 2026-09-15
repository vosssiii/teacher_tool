#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Prueft eine Word- oder PowerPoint-Vorlage auf die noetigen Formatvorlagen
(KONZEPT.md Kap. 6.1, offener Punkt 3).

Der Gedanke dahinter: Die Schule liefert eine Vorlage, die niemand fuer
dieses Plugin gebaut hat. Sie enthaelt also keine Formatvorlagen mit
unseren Namen, sondern die ganz gewoehnlichen von Word - "Standard",
"Ueberschrift 1" und so weiter, je nach Sprache der Installation auch auf
Englisch.

Deshalb wird nicht nach festen Namen gesucht, sondern je *Zweck* nach der
ersten brauchbaren Formatvorlage: erst unsere eigene UA-Vorlage, dann die
ueblichen Word-Namen in beiden Sprachen. Gefunden wird eine Zuordnung
Zweck -> tatsaechlicher Name, mit der md_zu_docx.py spaeter arbeitet.

Aufruf:
    python vorlage_pruefen.py --datei vorlage.docx
    python vorlage_pruefen.py --datei vorlage.pptx --json
"""

import argparse
import json
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# Zweck -> moegliche Formatvorlagen, beste zuerst.
# UA-* sind unsere eigenen; danach die Word-Namen auf Deutsch und Englisch.
ZWECKE = [
    ("titel", "Überschrift des Arbeitsblatts",
     ["UA-Titel", "Titel", "Title", "Überschrift 1", "Heading 1"], True),
    ("aufgabe", "Aufgabenüberschrift",
     ["UA-Aufgabe", "Überschrift 2", "Heading 2"], True),
    ("text", "Fließtext",
     ["UA-Text", "Standard", "Normal", "Textkörper", "Body Text"], True),
    ("material", "Materialien (M1, M2 …)",
     ["UA-Material", "Überschrift 3", "Heading 3"], False),
    ("loesung", "Lösungen (nur Lösungsversion)",
     ["UA-Loesung", "UA-Lösung", "Zitat", "Quote"], False),
    ("hinweis", "Hinweise für die Lehrkraft",
     ["UA-Hinweis", "Intensives Zitat", "Intense Quote", "Zitat", "Quote"], False),
    ("quelle", "Quellenangabe bei fremdem Material",
     ["UA-Quelle", "Beschriftung", "Caption"], True),
]


def word_pruefen(pfad):
    import docx
    dokument = docx.Document(str(pfad))
    vorhanden = {}
    for stil in dokument.styles:
        try:
            vorhanden[stil.name] = str(stil.type)
        except Exception:
            continue

    zuordnung, fehlend = {}, []
    for schluessel, beschreibung, kandidaten, noetig in ZWECKE:
        treffer = next((k for k in kandidaten if k in vorhanden), None)
        zuordnung[schluessel] = {
            "beschreibung": beschreibung,
            "gefunden": treffer,
            "eigene_formatvorlage": bool(treffer and treffer.startswith("UA-")),
            "noetig": noetig,
        }
        if noetig and not treffer:
            fehlend.append(beschreibung)

    return {
        "art": "word",
        "datei": str(pfad),
        "formatvorlagen_insgesamt": len(vorhanden),
        "zuordnung": zuordnung,
        "fehlend": fehlend,
        "brauchbar": not fehlend,
        "alle_namen": sorted(vorhanden),
    }


def powerpoint_pruefen(pfad):
    """Bei PowerPoint zaehlen nicht Formatvorlagen, sondern Folienlayouts.
    Gebraucht werden zwei: eines mit Titel, eines mit Titel und Inhalt."""
    from pptx import Presentation
    praesentation = Presentation(str(pfad))
    layouts = []
    for nummer, layout in enumerate(praesentation.slide_layouts):
        platzhalter = []
        for form in layout.placeholders:
            try:
                platzhalter.append(str(form.placeholder_format.type))
            except Exception:
                pass
        layouts.append({"nummer": nummer, "name": layout.name,
                        "platzhalter": platzhalter})

    mit_titel = [l for l in layouts if any("TITLE" in p for p in l["platzhalter"])]
    mit_inhalt = [l for l in layouts
                  if any("TITLE" in p for p in l["platzhalter"])
                  and any(("BODY" in p or "OBJECT" in p) for p in l["platzhalter"])]

    fehlend = []
    if not mit_titel:
        fehlend.append("ein Layout mit Titel")
    if not mit_inhalt:
        fehlend.append("ein Layout mit Titel und Inhalt")

    return {
        "art": "powerpoint",
        "datei": str(pfad),
        "layouts": layouts,
        "layout_titel": mit_titel[0]["nummer"] if mit_titel else None,
        "layout_titel_inhalt": mit_inhalt[0]["nummer"] if mit_inhalt else None,
        "fehlend": fehlend,
        "brauchbar": not fehlend,
    }


def bericht(daten):
    z = []
    if daten.get("fehler"):
        return "Die Vorlage konnte nicht gelesen werden: %s" % daten["fehler"]

    z.append("Vorlage geprüft: %s" % Path(daten["datei"]).name)
    z.append("=" * 60)

    if daten["art"] == "word":
        z.append("%d Formatvorlagen enthalten." % daten["formatvorlagen_insgesamt"])
        z.append("")
        for eintrag in daten["zuordnung"].values():
            if eintrag["gefunden"]:
                zusatz = "" if eintrag["eigene_formatvorlage"] else "  (Word-Standard)"
                z.append("  gefunden  %-38s → %s%s"
                         % (eintrag["beschreibung"], eintrag["gefunden"], zusatz))
            else:
                z.append("  fehlt     %-38s %s"
                         % (eintrag["beschreibung"],
                            "(nötig)" if eintrag["noetig"] else "(optional)"))
    else:
        z.append("%d Folienlayouts enthalten." % len(daten["layouts"]))
        z.append("  Layout mit Titel:          %s" % daten["layout_titel"])
        z.append("  Layout mit Titel + Inhalt: %s" % daten["layout_titel_inhalt"])

    z.append("")
    if daten["brauchbar"]:
        z.append("Ergebnis: Die Vorlage ist brauchbar.")
    else:
        z.append("Ergebnis: Die Vorlage reicht nicht. Es fehlt: %s"
                 % ", ".join(daten["fehlend"]))
        z.append("Das Plugin nimmt stattdessen seine eigene Standardvorlage.")
    return "\n".join(z)


def main():
    parser = argparse.ArgumentParser(
        description="Prüft eine Word- oder PowerPoint-Vorlage auf nötige Formatvorlagen")
    parser.add_argument("--datei", required=True, help="Die zu prüfende .docx oder .pptx")
    parser.add_argument("--json", action="store_true", help="Ergebnis als JSON ausgeben")
    argumente = parser.parse_args()

    pfad = Path(argumente.datei).expanduser()
    if not pfad.exists():
        daten = {"fehler": "Datei nicht gefunden: %s" % pfad, "brauchbar": False,
                 "datei": str(pfad)}
    else:
        try:
            if pfad.suffix.lower() == ".docx":
                daten = word_pruefen(pfad)
            elif pfad.suffix.lower() == ".pptx":
                daten = powerpoint_pruefen(pfad)
            else:
                daten = {"fehler": "Unbekannte Dateiart: %s" % pfad.suffix,
                         "brauchbar": False, "datei": str(pfad)}
        except BaseException as fehler:
            daten = {"fehler": "{}: {}".format(type(fehler).__name__, fehler),
                     "brauchbar": False, "datei": str(pfad)}

    if argumente.json:
        print(json.dumps(daten, indent=2, ensure_ascii=False, default=str))
    else:
        print(bericht(daten))
    # Kein Fehler-Exitcode: Eine unbrauchbare Vorlage ist ein Befund, kein
    # Absturz. Der Skill entscheidet, was daraus folgt.
    return 0


if __name__ == "__main__":
    try:
        ENDE = main()
    except BaseException:
        import traceback
        traceback.print_exc()
        ENDE = 1
    sys.exit(ENDE)
